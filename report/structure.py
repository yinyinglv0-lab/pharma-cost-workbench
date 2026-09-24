"""Read-only Module 1 template/content acceptance, independent of generation.

The specification is parsed from the actual DOCX, not from its Markdown copy.
Business levels are 1/2/3 even though the template has a preceding title level.
PDF extraction proves text/content order, never Word styles or PDF visual layout.
All public results are JSON-safe; no repository, model, RPA or scoring writes occur.
"""
from __future__ import annotations

import base64
from collections import Counter
from collections.abc import Mapping
import hashlib
from io import BytesIO
import re
import unicodedata

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from .registry import TEMPLATE_PATH

VERSION = "template-content-acceptance/2.0"
OFFICIAL_TEMPLATE_SHA256 = "188cbf06bd08df87acb1d9b4f33a0a13074a971249f14a9b0f1ca15544c509d6"
PROJECT_TEMPLATE_SHA256 = "6ff4239093db1f003744b8fa1c5644a4e2c2fd0a7630e0edb157b8f450971760"
NODE_IDS = ("1", "2", "2.1", "2.2", "3", "3.1", "3.1.1", "3.1.2", "3.2", "3.3",
            "4", "4.1", "4.2", "4.3", "5", "5.1", "5.2", "5.3", "6", "6.1", "6.2", "6.3", "6.4")
LEGACY_CHAPTERS = ("一、封面与基本信息", "二、总成本概览", "三、成本要素明细分析",
                   "四、重点产品专项分析", "五、对标分析", "六、总结与建议")
# Reviewed semantic equivalents only. Matching a numeric prefix alone is forbidden.
TITLE_EQUIVALENTS = {
    "2.2": ("2.2 成本结构与金额桥接",),
    "4": ("四、重点产品专项分析",),
    "4.1": ("4.1 近六个月单位成本趋势", "4.1 近六个月趋势与期间范围"),
    "4.2": ("4.2 专项问题与异常排查",),
    "4.3": ("4.3 原材料市场参考行情",),
    "5": ("五、对标分析",),
    "5.1": ("5.1 同期同规格差异",),
    "5.3": ("5.3 差异原因核查",),
    "6.1": ("6.1 本期成本管理亮点", "6.1 本期管理亮点", "6.1 本期经营观察"),
    "6.2": ("6.2 需关注问题",),
    "6.4": ("6.4 整改任务草稿",),
}
_NUMERIC_HEADING = re.compile(r"^([1-6](?:\.[1-9]\d*)+)\s*(.+)$")
_MAIN_HEADING = re.compile(r"^([一二三四五六])、(.+)$")
_PLACEHOLDER = re.compile(r"\{\{(.*?)\}\}", re.S)
_EMPTY = {"", "—", "-", "--", "n/a", "na", "null", "none", "nan", "inf", "infinity", "暂无", "无", "略",
          "待填写", "待补充", "待生成", "待完善", "未提供", "缺少数据", "暂无数据", "无数据", "数据缺失",
          "未提供数据", "见附录", "详见附录", "todo", "tbd"}
_MISSING = re.compile(r"缺少|缺失|缺月|未提供|未取得|未选择|暂无|不可得|不可比|无可用|无有效|无法|不适用|无定义|不计算|零分母|分母为零")
_PRODUCTS = ("银黄口服液", "板蓝根颗粒", "六味地黄胶囊")
# Minimum field families for business tables; words may occur in labels or units.
_TABLE_FIELDS = {
    "1": (("项目",), ("内容",)),
    "2.1": (("指标",), ("本期", "本月"), ("元", "盒")),
    "2.2": (("要素",), ("元",), ("占比",)),
    "3.1.1": (("原材料",), ("本期", "本月"), ("元/盒",)),
    "3.2": (("指标",), ("本期", "本月"), ("工时", "时薪", "人工",)),
    "3.3": (("费用类别",), ("本期", "本月"), ("元/盒",)),
    "4.1": (("月份",), ("成本", "材料"), ("元/盒",)),
    "4.3": (("药材", "原材料"), ("参考价", "本月价"), ("单位", "元/")),
    "5.1": (("月份", "对比维度"), ("一厂",), ("二厂",), ("元", "金额")),
    "6.3": (("建议",), ("责任",), ("优先",)),
    "6.4": (("任务",), ("责任",), ("截止",)),
}


def _norm(value):
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", str("" if value is None else value)))


def _number(text):
    value = unicodedata.normalize("NFKC", str(text)).strip()
    match = _MAIN_HEADING.fullmatch(value)
    if match:
        return str("一二三四五六".index(match[1]) + 1)
    match = _NUMERIC_HEADING.fullmatch(value)
    return match[1] if match else None


def _outline(paragraph):
    """Resolve direct outline or inherited paragraph-style outline, zero based."""
    properties = [paragraph._p.pPr]
    style, seen = paragraph.style, set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        properties.append(style.element.pPr)
        style = style.base_style
    for props in properties:
        node = props.find(qn("w:outlineLvl")) if props is not None else None
        if node is not None:
            value = int(node.get(qn("w:val")))
            return value if 0 <= value < 9 else None
    match = re.fullmatch(r"(?:heading|标题)\s*([1-9])", paragraph.style.name, re.I)
    return int(match[1]) - 1 if match else None


def template_structure(template_bytes=None):
    """Parse and verify the 23-node business contract from supplied/frozen DOCX bytes.

    A changed template cannot silently shrink the required denominator. The official
    hash is only a registered identity; the configured path is not called an original.
    """
    raw = TEMPLATE_PATH.read_bytes() if template_bytes is None else bytes(template_bytes)
    document = Document(BytesIO(raw))
    candidates = []
    for index, paragraph in enumerate(document.paragraphs):
        ident, outline = _number(paragraph.text), _outline(paragraph)
        if ident is not None and outline is not None:
            candidates.append((ident, paragraph.text.strip(), outline, index))
    if tuple(item[0] for item in candidates) != NODE_IDS:
        raise ValueError("DOCX模板业务标题必须完整且有序地包含23个节点；禁止缩小验收分母")
    base = candidates[0][2]
    nodes = []
    for ident, title, outline, position in candidates:
        level = ident.count(".") + 1
        if outline - base + 1 != level:
            raise ValueError("DOCX模板标题层级与业务编号不一致：" + ident)
        nodes.append({"id": ident, "title": title, "level": level,
                      "parent_id": ident.rsplit(".", 1)[0] if "." in ident else None,
                      "template_paragraph": position, "template_outline_level": outline,
                      "equivalent_titles": list(TITLE_EQUIVALENTS.get(ident, ()))})
    digest = hashlib.sha256(raw).hexdigest()
    return {"sha256": digest, "input": "configured_project_path" if template_bytes is None else "supplied_bytes",
            "path": str(TEMPLATE_PATH) if template_bytes is None else None,
            "identity": ("official_original" if digest == OFFICIAL_TEMPLATE_SHA256 else
                         "project_extended_copy" if digest == PROJECT_TEMPLATE_SHA256 else "other_template_bytes"),
            "official_original_registered_sha256": OFFICIAL_TEMPLATE_SHA256,
            "matches_official_original": digest == OFFICIAL_TEMPLATE_SHA256,
            "table_count": len(document.tables), "node_count": len(nodes),
            "level_counts": {str(level): sum(node["level"] == level for node in nodes) for level in (1, 2, 3)},
            "nodes": nodes}


def _match(text, spec):
    value, ident = _norm(text), _number(text)
    if ident is None:
        return None
    node = next((node for node in spec["nodes"] if node["id"] == ident), None)
    if node is None:
        return None
    for title in [node["title"], *node["equivalent_titles"]]:
        if "{{产品名称}}" in title:
            # The sole template parameter inside a business heading, not free fuzziness.
            prefix, suffix = _norm(title).split("{{产品名称}}")
            if value.startswith(prefix) and (not suffix or value.endswith(suffix)):
                product = value[len(prefix):len(value) - len(suffix) if suffix else None]
                if product and "{{" not in product and len(product) <= 80:
                    return ident
        elif value == _norm(title):
            return ident
    return None


def _residues(text):
    normalized = unicodedata.normalize("NFKC", str(text or ""))
    values = [re.sub(r"\s+", "", match) for match in _PLACEHOLDER.findall(normalized)]
    if ("{{" in normalized or "}}" in normalized) and not values:
        values.append("未闭合占位符")
    return values


def _strings(block):
    if block.get("kind") == "table":
        headers, rows = block.get("headers"), block.get("rows")
        if isinstance(headers, (list, tuple)):
            yield from (str(value) if value is not None else "" for value in headers)
        for row in rows if isinstance(rows, (list, tuple)) else []:
            if isinstance(row, (list, tuple)):
                yield from (str(value) if value is not None else "" for value in row)
        if block.get("note"):
            yield str(block["note"])
    else:
        yield str(block.get("text", block.get("caption", "")) or "")


def _nonempty(value):
    norm = _norm(value).strip("。；;，,：:").lower()
    return norm not in _EMPTY and bool(re.search(r"[\w\u3400-\u9fff]", norm)) and not _residues(value)


def _explanation(text):
    norm = _norm(text)
    if norm in ("无可用完整记录;见数据限制说明。", "缺少数据", "暂无数据"):
        return False
    return bool(_MISSING.search(norm) and len(re.findall(r"[\w\u3400-\u9fff]", norm)) >= 8)


def _prose(text, spec):
    norm = _norm(text)
    if not _nonempty(text) or _match(text, spec):
        return False
    if norm == "无可用完整记录;见数据限制说明。" or re.match(r"^(?:本节|此处)?(?:待填写|待补充|待生成|待完善|TODO|TBD)", norm, re.I):
        return False
    return _explanation(text) or len(re.findall(r"[\w\u3400-\u9fff]", norm)) >= 4


def _chart_bytes(value):
    if not isinstance(value, Mapping) or not value.get("base64"):
        return None
    try:
        raw = base64.b64decode(value["base64"], validate=True)
        if not raw or value.get("sha256") != hashlib.sha256(raw).hexdigest():
            return None
        from PIL import Image
        with Image.open(BytesIO(raw)) as image:
            image.verify()
        return raw
    except (ValueError, TypeError, OSError):
        return None


def _content(block, node_id, spec, charts, *, extracted_images=False):
    kind, errors = block.get("kind"), []
    text = "\n".join(_strings(block))
    if _residues(text):
        errors.append("unresolved_placeholder")
    if kind == "paragraph":
        valid = _prose(text, spec)
        return valid and not errors, "missing_data_explanation" if _explanation(text) else "text", errors
    if kind == "chart":
        valid = bool((extracted_images and block.get("image_sha256") and block.get("pixel_sha256")) or
                     _chart_bytes((charts or {}).get(block.get("name"))))
        if not valid:
            errors.append("unverified_or_empty_chart")
        return valid and not errors, "chart", errors
    if kind != "table":
        return False, str(kind), errors
    headers, rows = block.get("headers", []), block.get("rows", [])
    if not isinstance(headers, (list, tuple)) or not headers or not all(_nonempty(value) for value in headers):
        errors.append("empty_table_headers")
    if not isinstance(headers, (list, tuple)):
        headers = []
    if not isinstance(rows, (list, tuple)) or not rows:
        errors.append("empty_table")
        rows = []
    well_formed = bool(rows) and all(isinstance(row, (list, tuple)) and len(row) == len(headers) for row in rows)
    if rows and not well_formed:
        errors.append("invalid_table_shape")
    # Row labels alone (or all dashes) are not effective data. Specific missing-data
    # explanations may stand in for unavailable values, but never an empty table.
    numeric_table = node_id in ("2.1", "2.2", "3.1.1", "3.2", "3.3", "4.1", "4.3", "5.1")
    value_columns = [index for index, header in enumerate(headers) if index > 0 and
                     (not numeric_table or re.search(r"本期|本月|前期|上月|元|盒|价|金额|成本|差|率|占比|贡献|一厂|二厂", str(header)))]
    def valid_value(value):
        return _nonempty(value) and (not numeric_table or bool(re.fullmatch(r"[+-]?\d[\d,.]*(?:%|元(?:/盒)?)?", _norm(value))))
    data_rows = sum(any(valid_value(row[index]) for index in value_columns if index < len(row)) or
                    _explanation(" ".join(map(str, row))) for row in rows if isinstance(row, (list, tuple)))
    explanation = _explanation(block.get("note", block.get("_validation_note", "")))
    if rows and not data_rows and not explanation:
        errors.append("empty_table_values")
    field_text = _norm(" ".join(map(str, headers)) + " " + " ".join(str(row[0]) for row in rows if isinstance(row, (list, tuple)) and row))
    for family in _TABLE_FIELDS.get(node_id, ()):
        if not any(_norm(term) in field_text for term in family):
            errors.append("missing_table_field:" + "/".join(family))
    valid = bool(headers and rows and well_formed and (data_rows or explanation))
    return valid and not errors, "table", errors


def _period_months(period):
    if period is None:
        return []
    value = period.get("months", []) if isinstance(period, Mapping) else period
    if isinstance(value, str):
        quarter = re.fullmatch(r"(\d{4})[Qq]([1-4])", value)
        if quarter:
            return [f"{quarter[1]}-{month:02d}" for month in range((int(quarter[2]) - 1) * 3 + 1, int(quarter[2]) * 3 + 1)]
        value = [value]
    if not isinstance(value, (list, tuple)) or not value or any(not re.fullmatch(r"\d{4}-(?:0[1-9]|1[0-2])", str(month)) for month in value):
        raise ValueError("period须提供YYYY-MM、YYYYQn或包含完整months列表的对象")
    return sorted(set(map(str, value)))


def _months_in(text):
    value = unicodedata.normalize("NFKC", text)
    months = {f"{year}-{int(month):02d}" for year, month in re.findall(r"(\d{4})[-年](0?[1-9]|1[0-2])(?:月|(?=[^\d]|$))", value)}
    for year, quarter in re.findall(r"(\d{4})\s*[Qq]([1-4])", value):
        months.update(_period_months(year + "Q" + quarter))
    return sorted(months)


def _scope(blocks, entries, spec, product, period):
    expected_months = _period_months(period)
    result = {"status": "not_requested" if product is None and period is None else "checked",
              "product": {"expected": product, "passed": None},
              "period": {"expected_months": expected_months, "passed": None},
              "limitations": "检查正文封面元信息及显式场景冲突；不能证明每项数值和经营归因真实。"}
    cover = next((entry for entry in entries if entry["id"] == "1"), None)
    if cover:
        end = next((entry["index"] for entry in entries if entry["index"] > cover["index"]), len(blocks))
        cover_blocks = blocks[cover["index"] + 1:end]
    else:
        cover_blocks = []
    cover_text = "\n".join(value for block in cover_blocks for value in _strings(block))
    product_values, period_values = [], []
    for block in cover_blocks:
        if block.get("kind") == "table":
            rows = block.get("rows", [])
            for row in rows if isinstance(rows, (list, tuple)) else []:
                if not isinstance(row, (list, tuple)) or len(row) < 2:
                    continue
                label, value = _norm(row[0]), " ".join(map(str, row[1:]))
                if label in ("产品/规格", "产品名称", "分析产品", "产品"):
                    product_values.append(value)
                if label in ("期间", "分析周期", "分析月份", "月份"):
                    period_values.append(value)
    if product is not None:
        product_view = "\n".join(product_values) if product_values else cover_text
        result["product"].update(observed_values=product_values,
                                  passed=bool(_norm(product) and _norm(product) in _norm(product_view)))
        foreign = [name for name in _PRODUCTS if name != product and name in product_view]
        if foreign:
            result["product"].update(passed=False, conflicting_products=foreign)
        # A correct cover must not conceal another scenario pasted into the body.
        # Only business prose is checked; source catalogues/appendices are excluded.
        body_conflicts, in_appendix = [], False
        for block in blocks:
            text = str(block.get("text", ""))
            if block.get("kind") == "heading" and text.strip().startswith("附录"):
                in_appendix = True
            if in_appendix or block.get("kind") != "paragraph" or re.search(r"不适用|非本产品|仅作对比|不可套用", text):
                continue
            names = [name for name in _PRODUCTS if name != product and name in text]
            if names:
                body_conflicts.append({"products": names, "text": text[:160]})
        if body_conflicts:
            result["product"].update(passed=False, body_conflicts=body_conflicts)
    if period is not None:
        period_view = "\n".join(period_values) if period_values else cover_text
        observed = _months_in(period_view)
        result["period"].update(observed_months=observed, passed=observed == expected_months)
    result["passed"] = all(result[key]["passed"] is not False for key in ("product", "period"))
    return result


def _validate(blocks, spec, *, product=None, period=None, charts=None, source="blocks", hierarchy=True):
    blocks = [dict(block) for block in blocks]
    errors, entries, excluded, stack = [], [], [], {}
    appendix, toc = False, False
    previous_rank = -1
    for index, block in enumerate(blocks):
        kind, text = block.get("kind"), str(block.get("text", ""))
        if kind == "heading" and not text.strip():
            errors.append({"code": "empty_heading", "block_index": index})
            continue
        if kind == "heading" and re.match(r"^附录(?:[一二三四五六\d:：\s]|$)", text.strip()):
            appendix = True
        if block.get("role") == "toc" or kind == "toc" or _norm(text) == "目录":
            toc = True
        elif toc and kind == "heading" and _match(text, spec) == "1":
            # A real cover has body content before the next business heading.
            following = []
            for candidate in blocks[index + 1:]:
                if candidate.get("kind") == "heading":
                    break
                following.append(candidate)
            toc = not any(candidate.get("kind") == "table" or
                          candidate.get("kind") == "paragraph" and _prose(candidate.get("text", ""), spec)
                          for candidate in following)
        if appendix or toc or block.get("role") in ("toc", "header", "footer"):
            excluded.append(index)
            continue
        if kind != "heading":
            continue
        ident = _match(text, spec)
        raw_level = block.get("level", 1) if hierarchy else None
        level = raw_level if type(raw_level) is int and 1 <= raw_level <= 9 else None
        if ident is None:
            if _number(text) is not None:
                errors.append({"code": "unregistered_heading", "block_index": index, "text": text})
            if hierarchy and level:
                stack = {depth: value for depth, value in stack.items() if depth < level}
                stack[level] = None
            continue
        node = next(node for node in spec["nodes"] if node["id"] == ident)
        rank, issues = NODE_IDS.index(ident), []
        if rank <= previous_rank:
            issues.append("out_of_order")
        previous_rank = max(previous_rank, rank)
        if hierarchy:
            if level != node["level"]:
                issues.append("wrong_level")
            actual_parent = stack.get((level or 1) - 1)
            if actual_parent != node["parent_id"]:
                issues.append("wrong_parent")
            stack = {depth: value for depth, value in stack.items() if depth < (level or 1)}
            stack[level or 1] = ident
        entries.append({"id": ident, "index": index, "actual_level": level, "errors": issues,
                        "text": text, "location": block.get("location", {"block_index": index})})
    excluded_set = set(excluded)
    occurrences = Counter(entry["id"] for entry in entries)
    content_checks = []
    for index, block in enumerate(blocks):
        if index in excluded_set or block.get("kind") not in ("paragraph", "table", "chart"):
            continue
        owner = next((entry["id"] for entry in reversed(entries) if entry["index"] < index), None)
        if owner is None:
            continue
        valid, category, issues = _content(block, owner, spec, charts, extracted_images=source == "docx")
        content_checks.append({"block_index": index, "node_id": owner, "kind": category,
                               "valid": valid, "errors": issues, "location": block.get("location", {"block_index": index})})
    node_results = []
    for node in spec["nodes"]:
        found = [entry for entry in entries if entry["id"] == node["id"]]
        issues = []
        content = []
        if not found:
            issues.append("missing_heading")
        else:
            entry = found[0]
            issues.extend(entry["errors"])
            if occurrences[node["id"]] != 1:
                issues.append("duplicate_heading")
            end = next((candidate["index"] for candidate in entries
                        if candidate["index"] > entry["index"] and candidate["id"].count(".") + 1 <= node["level"]), len(blocks))
            content = [item for item in content_checks if entry["index"] < item["block_index"] < end]
            if not any(item["valid"] for item in content):
                issues.append("empty_content")
            for item in content:
                issues.extend(item["errors"])
        content_valid = bool(found and any(item["valid"] for item in content) and not any(item["errors"] for item in content))
        own_valid = [item for item in content if item["node_id"] == node["id"] and item["valid"]]
        status = ("missing" if not found else "invalid" if any(item["errors"] for item in content) else
                  "empty" if not content_valid else "subsections" if not own_valid else
                  "missing_data_explanation" if all(item["kind"] == "missing_data_explanation" for item in own_valid) else "present")
        node_results.append({**node, "present": bool(found), "occurrences": len(found),
                             "locations": [entry["location"] for entry in found],
                             "actual_title": found[0]["text"] if found else None,
                             "actual_level": found[0]["actual_level"] if found else None,
                             "hierarchy_passed": None if not hierarchy else bool(found and not set(issues) & {"wrong_level", "wrong_parent"}),
                             "order_passed": bool(found and not set(issues) & {"out_of_order", "duplicate_heading"}),
                             "content_status": status, "content_valid": content_valid,
                             "content_blocks": content, "errors": list(dict.fromkeys(issues)), "passed": not issues})
    residues = [{"block_index": index, "placeholders": values} for index, block in enumerate(blocks)
                if (values := _residues("\n".join(_strings(block))))]
    if residues:
        errors.append({"code": "unresolved_placeholders", "locations": residues})
    scope = _scope(blocks, entries, spec, product, period)
    if not scope["passed"]:
        errors.append({"code": "scope_mismatch"})
    def metric(rows, key):
        matched = sum(bool(row[key]) for row in rows)
        return {"matched": matched, "expected": len(rows), "rate": matched / len(rows)}
    chapters = [row for row in node_results if row["level"] == 1]
    complete = [dict(row, structure_valid=row["present"] and row["order_passed"] and row["hierarchy_passed"] is not False)
                for row in node_results]
    return {"schema_version": VERSION, "source": source, "template": spec, "nodes": node_results,
            "metrics": {"main_chapters": metric(chapters, "present"), "full_nodes": metric(node_results, "present"),
                        "ordered_structure": metric(complete, "structure_valid"), "effective_content": metric(node_results, "content_valid")},
            "hierarchy": {"status": "checked" if hierarchy else "not_observable", "passed": all(row["hierarchy_passed"] for row in node_results) if hierarchy else None},
            "validation_scope": "semantic_levels_content_and_order" if hierarchy else "extracted_content_and_order_only",
            "format_properties_verified": False, "visual_review": "not_performed",
            "scope_checks": scope, "placeholder_residues": residues,
            "excluded_block_indices": excluded, "errors": errors,
            "passed": not errors and all(row["passed"] for row in node_results)}


def validate_blocks(blocks, *, template_bytes=None, product=None, period=None, charts=None):
    """Validate frozen semantic blocks; optional product/period bind the cover scope.

    Pass payload['charts'] to verify chart assets. A caption/name alone never counts
    as a valid chart. Parent chapters may carry content through their subsections.
    """
    return _validate(blocks, template_structure(template_bytes), product=product, period=period, charts=charts)


def _docx_blocks(raw):
    document = Document(BytesIO(raw))
    blocks, image_hashes, image_pixels = [], [], []
    for index, element in enumerate(document.element.body):
        location = {"body_index": index}
        if element.tag == qn("w:p"):
            paragraph = Paragraph(element, document)
            level = _outline(paragraph)
            style = paragraph.style.name.lower()
            toc = style.startswith("toc") or style.startswith("目录") or bool(element.xpath('.//w:hyperlink[starts-with(@w:anchor, "_Toc")]'))
            block = {"kind": "heading" if level is not None else "paragraph", "text": paragraph.text,
                     "level": level + 1 if level is not None else None, "location": location}
            if toc:
                block["role"] = "toc"
            blocks.append(block)
            for blip in element.xpath(".//a:blip"):
                relationship = blip.get(qn("r:embed"))
                if relationship and relationship in document.part.related_parts:
                    image_raw = document.part.related_parts[relationship].blob
                    digest, pixels = hashlib.sha256(image_raw).hexdigest(), _pixel_hash(image_raw)
                    image_hashes.append(digest)
                    image_pixels.append(pixels)
                    blocks.append({"kind": "chart", "image_sha256": digest, "pixel_sha256": pixels, "location": location})
        elif element.tag == qn("w:tbl"):
            table = Table(element, document)
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            blocks.append({"kind": "table", "headers": rows[0] if rows else [], "rows": rows[1:], "location": location})
        elif element.tag == qn("w:sdt"):
            # TOCs/other structured fields are not body section evidence.
            text = "\n".join(element.xpath(".//w:t/text()"))
            blocks.append({"kind": "toc", "role": "toc", "text": text, "location": location})
    # Exported table notes are adjacent paragraphs, not table XML. Bind only the
    # immediate note for validity (not a second copy for same-source text matching).
    for index, block in enumerate(blocks[:-1]):
        following = blocks[index + 1]
        if block.get("kind") == "table" and following.get("kind") == "paragraph" and _explanation(following.get("text", "")):
            block["_validation_note"] = following["text"]
    # Audit all visible XML parts, including nested tables/header/footer. It is
    # intentionally separate from section coverage so appendix residue still fails.
    extra_text = []
    for part in document.part.package.parts:
        if part.partname.startswith("/word/") and hasattr(part, "element"):
            extra_text.extend(part.element.xpath(".//w:t/text()"))
    return blocks, {"image_hashes": image_hashes, "image_pixels": image_pixels, "all_text": "\n".join(extra_text),
                    'metadata_subject': document.core_properties.subject}


def _pdf_blocks(raw, spec):
    from pypdf import PdfReader
    reader = PdfReader(BytesIO(raw))
    blocks, images, all_text = [], [], []
    for page_number, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        all_text.append(text)
        lines, index = text.splitlines(), 0
        while index < len(lines):
            line = lines[index].strip()
            location = {"page": page_number, "line": index + 1}
            index += 1
            if not line or _norm(line) == "中药一厂·成本分析报告·内部" or re.fullmatch(r"第\s*\d+\s*页", line):
                continue
            if re.fullmatch(r"CB-[A-Z0-9]+\s*\|\s*(?:待审核|已审核签发)(?:\s*\|\s*[a-f0-9]{12})?", line):
                continue
            ident = _match(line, spec)
            if not ident and _number(line) and index < len(lines):
                joined = line + lines[index].strip()
                if _match(joined, spec):
                    line, ident = joined, _match(joined, spec)
                    index += 1
            heading = bool(ident or re.match(r"^附录[一二三四五六\d:：]", line))
            blocks.append({"kind": "heading" if heading else "paragraph", "text": line, "location": location})
        for image in page.images:
            images.append({"page": page_number, "pixel_sha256": _pixel_hash(image.data)})
    return blocks, {"pages": len(reader.pages), "images": images, "all_text": "\n".join(all_text),
                    'metadata_subject': (reader.metadata or {}).get('/Subject', '')}


def _pixel_hash(raw):
    try:
        from PIL import Image
        with Image.open(BytesIO(raw)) as image:
            rgb = image.convert("RGB")
            return hashlib.sha256(str(rgb.size).encode("ascii") + b":" + rgb.tobytes()).hexdigest()
    except (ValueError, OSError):
        return None


def _sections(blocks, spec):
    result, current, appendix = {}, None, False
    for block in blocks:
        text = block.get("text", "")
        if block.get("kind") == "heading":
            if re.match(r"^附录[一二三四五六\d:：]", text.strip()):
                appendix = True
            if not appendix and block.get("role") != "toc":
                current = _match(text, spec)
                if current:
                    result.setdefault(current, [])
            else:
                current = None
            continue
        if current and not appendix and block.get("role") != "toc":
            result[current].extend(value for value in _strings(block) if _norm(value))
    return result


def _compare_text(reference, target, spec):
    source_sections, target_sections = _sections(reference, spec), _sections(target, spec)
    details = []
    for ident in NODE_IDS:
        atoms = source_sections.get(ident, [])
        raw_view = unicodedata.normalize("NFKC", "\n".join(target_sections.get(ident, [])))
        positions = [index for index, char in enumerate(raw_view) if not char.isspace()]
        haystack = "".join(raw_view[index] for index in positions)
        cursor, missing = 0, []
        for atom in atoms:
            needle = _norm(atom)
            found = haystack.find(needle, cursor)
            if re.fullmatch(r"[+-]?\d[\d,.]*%?", needle):
                # A cell 3.50 must not match 13.50, -3.50 or 3.500. Keep the
                # original whitespace boundaries between adjacent exported cells.
                while found >= 0:
                    start, end = positions[found], positions[found + len(needle) - 1] + 1
                    before, after = raw_view[start - 1:start] if start else "", raw_view[end:end + 1]
                    if not (before and before in "0123456789.,+-") and not (after and after in "0123456789.,%"):
                        break
                    found = haystack.find(needle, found + 1)
            if found < 0:
                missing.append(atom[:160])
            else:
                cursor = found + len(needle)
        details.append({"id": ident, "expected_atoms": len(atoms), "missing_or_reordered": missing, "passed": not missing})
    return {"passed": all(row["passed"] for row in details), "nodes": details,
            "normalization": "NFKC_and_whitespace_only", "scope": "each_required_section_body_in_order"}


def _bind_source_content(result, reference, comparison, basis):
    for output_node, source_node, text_check in zip(result["nodes"], reference["nodes"], comparison["nodes"]):
        output_node["source_content_basis"] = basis
        output_node["source_content_status"] = source_node["content_status"]
        if not source_node["content_valid"] or not text_check["passed"]:
            output_node["content_valid"] = False
            output_node["content_status"] = "source_or_export_content_invalid"
            output_node["errors"].append("source_content_invalid" if not source_node["content_valid"] else "source_content_mismatch")
            output_node["passed"] = False
    metric = result["metrics"]["effective_content"]
    metric["matched"] = sum(row["content_valid"] for row in result["nodes"])
    metric["rate"] = metric["matched"] / metric["expected"]
    result["passed"] = result["passed"] and all(row["passed"] for row in result["nodes"])


def _failed_export(spec, format, reason):
    result = _validate([], spec, source=format, hierarchy=format == "docx")
    result["errors"].append({"code": reason})
    return result


def export_checks(docx, pdf, *, blocks=None, template_bytes=None, product=None, period=None, charts=None, frozen_hash=None):
    """Check both exports at the same source-content level, retaining legacy keys.

    With blocks, compare each export against the frozen block content. Without them,
    compare PDF against DOCX and explicitly do not claim frozen-source verification.
    Absent/unreadable formats still contribute all 23 required nodes to the denominator.
    """
    spec = template_structure(template_bytes)
    structure, parsed, info = {}, {}, {}
    source_blocks = [dict(block) for block in blocks] if blocks is not None else None
    source_validation = (_validate(source_blocks, spec, product=product, period=period, charts=charts)
                         if source_blocks is not None else None)
    for format, raw, parser in (("docx", docx, _docx_blocks), ("pdf", pdf, lambda raw: _pdf_blocks(raw, spec))):
        if not raw:
            structure[format] = _failed_export(spec, format, "export_missing")
            parsed[format], info[format] = [], {"all_text": ""}
            continue
        try:
            parsed[format], info[format] = parser(raw)
            structure[format] = _validate(parsed[format], spec, product=product, period=period,
                                           source=format, hierarchy=format == "docx")
            residues = _residues(info[format]["all_text"])
            if residues:
                structure[format]["errors"].append({"code": "export_part_placeholders", "placeholders": residues})
                structure[format]["passed"] = False
        except Exception as exc:
            # Fail closed without leaking third-party exception content into artifacts.
            structure[format] = _failed_export(spec, format, "export_unreadable:" + type(exc).__name__)
            parsed[format], info[format] = [], {"all_text": ""}
    reference = source_blocks if source_blocks is not None else parsed["docx"]
    comparisons = {format: _compare_text(reference, parsed[format], spec) for format in ("docx", "pdf")}
    expected_charts = [block for block in reference if block.get("kind") == "chart"]
    docx_hashes = info["docx"].get("image_hashes", [])
    pdf_pixels = [image["pixel_sha256"] for image in info["pdf"].get("images", [])]
    chart_checks = {"expected": len(expected_charts), "docx": len(docx_hashes), "pdf": len(pdf_pixels),
                    "mode": "asset_hash_and_pixels" if charts is not None else "docx_pdf_pixels_and_caption",
                    "passed": len(docx_hashes) == len(expected_charts) == len(pdf_pixels) and
                              None not in pdf_pixels and Counter(info["docx"].get("image_pixels", [])) == Counter(pdf_pixels)}
    if charts is not None:
        expected_raw = [_chart_bytes(charts.get(block.get("name"))) for block in expected_charts]
        valid_assets = all(value is not None for value in expected_raw)
        chart_checks["passed"] = bool(chart_checks["passed"] and valid_assets and
                                         Counter(hashlib.sha256(raw).hexdigest() for raw in expected_raw if raw is not None) == Counter(docx_hashes) and
                                         Counter(_pixel_hash(raw) for raw in expected_raw if raw is not None) == Counter(pdf_pixels))
    # Both content metrics share the source semantics. Extracted PDF header words
    # cannot make an empty source table valid; PDF heading styles remain unclaimed.
    if source_validation is not None:
        _bind_source_content(structure["docx"], source_validation, comparisons["docx"], "supplied_blocks")
    _bind_source_content(structure["pdf"], source_validation if source_validation is not None else structure["docx"],
                         comparisons["pdf"], "supplied_blocks" if source_validation is not None else "docx")
    markers = {format: re.findall(r"冻结报告SHA256[:：]([a-f0-9]{64})", _norm(info[format]["all_text"]))
               for format in ("docx", "pdf")}
    metadata_hashes = {format: re.findall(r'^冻结报告\s+([a-f0-9]{64})$', info[format].get('metadata_subject', ''))
                       for format in ('docx', 'pdf')}
    # Renderer 3 keeps the full digest in document metadata / audit sidecar.
    # When a visible legacy marker exists, it must agree with metadata as well.
    observed = {format: markers[format] or metadata_hashes[format] for format in markers}
    hash_checked = frozen_hash is not None or any(observed.values())
    marker_consistent = all(not markers[fmt] or not metadata_hashes[fmt] or markers[fmt] == metadata_hashes[fmt] for fmt in markers)
    hash_passed = ((all(observed[format] == [frozen_hash] for format in observed) if frozen_hash is not None else
                    bool(observed['docx'] and observed['docx'] == observed['pdf'])) and marker_consistent) if hash_checked else None
    same_source = {"mode": "supplied_blocks" if source_blocks is not None else "docx_reference",
                   "frozen_source_supplied": source_blocks is not None, "text": comparisons, "charts": chart_checks,
                   "frozen_hash": {"expected": frozen_hash, "observed": observed, 'visible_markers': markers,
                                    'metadata_hashes': metadata_hashes, "passed": hash_passed,
                                   "status": "checked_export_marker_or_metadata" if hash_checked else "not_supplied"},
                   "passed": all(row["passed"] for row in comparisons.values()) and chart_checks["passed"] and hash_passed is not False,
                   "limitations": "文本、引文、表格值与图像核验不代替视觉审阅或数据真实性/引用支持性人工评判。"}
    missing = {format: [title for index, title in enumerate(LEGACY_CHAPTERS, 1)
                        if not next(row for row in structure[format]["nodes"] if row["id"] == str(index))["present"]]
               for format in ("docx", "pdf")}
    return {"chapters_expected": list(LEGACY_CHAPTERS), "missing_chapters": missing,
            "chapter_completeness": {format: (6 - len(rows)) / 6 for format, rows in missing.items()},
            "placeholders_remaining": _residues(info["docx"]["all_text"]),
            "pdf_pages": info["pdf"].get("pages", 0), "charts_docx": len(docx_hashes), "charts_pdf": len(pdf_pixels),
            "task_draft_boundary_present": "未送达" in info["pdf"]["all_text"] and "未发送" in info["pdf"]["all_text"],
            "local_simulation_approval_present": "本机模拟验收" in info["pdf"]["all_text"],
            "full_node_completeness": {format: value["metrics"]["full_nodes"]["rate"] for format, value in structure.items()},
            "content_completeness": {format: value["metrics"]["effective_content"]["rate"] for format, value in structure.items()},
            "structure": structure, "source_validation": source_validation, "same_source": same_source,
            "counting_basis": {"internal_metric": True, "nodes_per_format": len(spec["nodes"]), "formats_expected": 2,
                               "expected_nodes_this_report": len(spec["nodes"]) * 2,
                               "three_started_scenarios_two_formats": len(spec["nodes"]) * 6,
                               "note": "138仅为当前23节点×三场景×两格式内部口径，不是官方评分公式；缺失格式不减分母。"},
            "passed": all(value["passed"] for value in structure.values()) and same_source["passed"] and
                      (source_validation is None or source_validation["passed"])}
