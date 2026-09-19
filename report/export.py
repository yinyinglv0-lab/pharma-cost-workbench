"""Template-based DOCX and ReportLab PDF from the same sealed report payload.

PDF uses an embedded TrueType CJK font. Word/COM, office automation, network
fetching, shell converters and recomputation from live source tables are absent.
"""
from __future__ import annotations

import base64
import hashlib
from importlib.metadata import version, PackageNotFoundError
import io
import os
from pathlib import Path
from threading import RLock
from xml.sax.saxutils import escape

FONT_ENV = "REPORT_CJK_FONT"
_RENDER_LOCK = RLock()


def _package(name):
    try:
        return version(name)
    except PackageNotFoundError:
        return "unavailable"


def font_descriptor(required_text=''):
    """Select an embeddable font covering the actual frozen report evidence."""
    try:
        from reportlab.pdfbase.ttfonts import TTFont
    except ImportError as exc:
        raise RuntimeError("PDF依赖reportlab未安装，请由部署环境安装reportlab") from exc
    candidates = [os.environ.get(FONT_ENV),
                  "C:/Windows/Fonts/HarmonyOS_Sans_SC_Regular.ttf",
                  "C:/Windows/Fonts/Deng.ttf", "C:/Windows/Fonts/msyh.ttc",
                  "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
                  "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
                  "/usr/share/fonts/truetype/noto/NotoSansSC-Regular.ttf",
                  "/usr/share/fonts/truetype/noto/NotoSansSC-VF.ttf",
                  "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
                  "/System/Library/Fonts/PingFang.ttc"]
    failures = []
    for value in dict.fromkeys(value for value in candidates if value):
        path = Path(value)
        if not path.is_file():
            continue
        try:
            font = TTFont("CJKProbe", str(path), subfontIndex=0)
            chars = set("成本报告材料制造费用分析来源月份季度专题盒采购实物耗用未送达" + required_text)
            if not all(char.isspace() or ord(char) in font.face.charToGlyph for char in chars):
                raise ValueError("字体缺少所需中文字符")
            return {"file": path.name, "path": str(path.resolve()),
                    "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                    "family": font.face.familyName.decode("utf-8", errors="replace") if isinstance(font.face.familyName, bytes) else str(font.face.familyName),
                    "subfont_index": 0, "embedded": True}
        except Exception as exc:
            failures.append(path.name + ":" + type(exc).__name__)
    raise RuntimeError("未找到ReportLab可嵌入的CJK TrueType字体；设置REPORT_CJK_FONT指向已授权的中文TTF或TrueType TTC。"
                       + "；".join(failures))


def renderer_versions(font):
    from .model import RENDERER_VERSION
    return {"version": RENDERER_VERSION, "docx": "python-docx/" + _package("python-docx"),
            "pdf": "reportlab/" + _package("reportlab"), "charts": "matplotlib/" + _package("matplotlib"),
            "font": font, "layout": "A4-six-sections/1.0"}


def _font_path(descriptor):
    candidates = [descriptor["path"], os.environ.get(FONT_ENV)]
    for value in candidates:
        if value and Path(value).is_file() and hashlib.sha256(Path(value).read_bytes()).hexdigest() == descriptor["sha256"]:
            return str(Path(value))
    raise ValueError("冻结报告所用字体已变化或缺失；请恢复相同哈希字体后重放")


def make_charts(facts, font):
    """Create PNGs once. The two documents embed exactly these saved bytes."""
    from matplotlib.figure import Figure
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.font_manager import FontProperties
    prop = FontProperties(fname=_font_path(font))
    charts = {}
    colors = ["#23678C", "#B57C42", "#358B80"]

    def save(figure, name, caption):
        for axis in figure.axes:
            for text in [axis.title, axis.xaxis.label, axis.yaxis.label, *axis.get_xticklabels(), *axis.get_yticklabels(), *axis.texts]:
                text.set_fontproperties(prop)
        buffer = io.BytesIO()
        FigureCanvasAgg(figure).print_png(buffer)
        raw = buffer.getvalue()
        charts[name] = {"base64": base64.b64encode(raw).decode("ascii"),
                        "sha256": hashlib.sha256(raw).hexdigest(), "caption": caption,
                        "width": 1080, "height": 430}
        figure.clear()

    with _RENDER_LOCK:
        figure = Figure(figsize=(9, 3.58), dpi=120, constrained_layout=True)
        axis = figure.subplots()
        trend = facts["trend"]
        axis.plot([row["month"] for row in trend], [row.get("unit_cost", float("nan")) for row in trend],
                  marker="o", color=colors[0], linewidth=2)
        axis.set_title("近六个月单位成本趋势（缺月保留断点）")
        axis.set_ylabel("元/盒")
        axis.grid(axis="y", alpha=.2)
        save(figure, "trend", "单位成本趋势截至分析期末；未提供月份不插值。")
        figure = Figure(figsize=(9, 3.58), dpi=120, constrained_layout=True)
        axis = figure.subplots()
        amounts = [row["amount"] for row in facts["elements"].values()]
        if sum(amounts) > 0:
            axis.pie(amounts, labels=["材料", "人工", "制造费用"], colors=colors, autopct="%.1f%%", startangle=90,
                     textprops={"fontproperties": prop})
        else:
            axis.text(.5, .5, "本期三要素金额均为零，结构占比无定义", ha="center", va="center", transform=axis.transAxes)
            axis.set_axis_off()
        axis.set_title("本期成本金额结构")
        save(figure, "structure", "结构占比按金额计算；零总额时不计算占比。")
        figure = Figure(figsize=(9, 3.58), dpi=120, constrained_layout=True)
        axis = figure.subplots()
        change = facts["amount_change"]
        if change.get("总变动额") is None:
            axis.text(.5, .5, "缺少完整可比前期，不绘制金额变化桥接", ha="center", va="center", transform=axis.transAxes)
            axis.set_axis_off()
        else:
            base = change["上月总成本"]
            axis.bar(0, base, color=colors[0])
            for index, key in enumerate(("材料", "人工", "制费"), 1):
                delta = change[key + "变动额"]
                after = base + delta
                axis.bar(index, abs(delta), bottom=min(base, after), color="#B86A4C" if delta >= 0 else colors[2])
                axis.annotate(f"{delta:+,.2f}", (index, max(base, after)), xytext=(0, 7), textcoords="offset points", ha="center", fontsize=8)
                base = after
            axis.bar(4, change["本月总成本"], color=colors[0])
            axis.set_xticks(range(5), ["完整前期", "材料变动", "人工变动", "制费变动", "本期"])
            axis.set_ylabel("元")
            axis.ticklabel_format(axis="y", style="plain")
            axis.grid(axis="y", alpha=.2)
            axis.margins(y=.2)
        axis.set_title("总成本金额桥接（含产量与单位成本共同影响）")
        save(figure, "waterfall", "瀑布图直接使用统一金额变动合同；零净变动时仍保留抵消要素。")
    return charts


def _verify_chart(chart):
    raw = base64.b64decode(chart["base64"], validate=True)
    if hashlib.sha256(raw).hexdigest() != chart["sha256"]:
        raise ValueError("图表字节哈希不一致")
    return raw


def _check_glyphs(payload, path):
    from reportlab.pdfbase.ttfonts import TTFont
    font = TTFont("ReportGlyphProbe", path, subfontIndex=0)
    text = "".join(block.get("text", "") + block.get("caption", "") + "".join(block.get("headers", []))
                   + "".join("".join(row) for row in block.get("rows", [])) for block in payload["blocks"])
    missing = sorted({char for char in text if not char.isspace() and ord(char) not in font.face.charToGlyph})
    if missing:
        raise ValueError("报告字体缺少字符：" + "".join(missing[:40]))


def _docx_font(run, size=10, bold=None, family="Microsoft YaHei"):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    run.font.name = family
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), family)


def render_docx(payload):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from .model import verify_payload
    verify_payload(payload)
    family = payload["versions"]["renderer"]["font"]["family"]
    def apply_font(run, size=10, bold=None):
        _docx_font(run, size, bold, family=family)
    doc = Document(io.BytesIO(base64.b64decode(payload["template_base64"])))
    from docx.enum.style import WD_STYLE_TYPE
    for name, kind in (("Title", WD_STYLE_TYPE.PARAGRAPH), ("Table Grid", WD_STYLE_TYPE.TABLE)):
        if name not in doc.styles:
            style = doc.styles.add_style(name, kind)
            style.base_style = doc.styles["Normal" if kind == WD_STYLE_TYPE.PARAGRAPH else "Normal Table"]
    # Retain the template package/styles, replacing its generic solution preface
    # and empty distribution lists with the frozen business report sections.
    for child in list(doc._element.body):
        if child.tag != qn("w:sectPr"):
            doc._element.body.remove(child)
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.8)
        section.left_margin, section.right_margin = Cm(1.5), Cm(1.5)
        section.header_distance, section.footer_distance = Cm(.8), Cm(.8)
        section.different_first_page_header_footer = False
        for header in (section.header, section.first_page_header, section.even_page_header):
            for child in list(header._element):
                header._element.remove(child)
            paragraph = header.add_paragraph("中药一厂 · 成本分析报告 · 内部")
            for run in paragraph.runs:
                apply_font(run, 8)
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            for child in list(footer._element):
                footer._element.remove(child)
            review_label = "已审核签发" if payload.get("review_status") == "approved" else "待审核"
            paragraph = footer.add_paragraph(payload["report_id"] + "  |  " + review_label + "  |  第 ")
            fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
            paragraph._p.append(fld)
            paragraph.add_run(" 页")
            for run in paragraph.runs:
                apply_font(run, 8)
    for block in payload["blocks"]:
        kind = block["kind"]
        if kind in ("title", "heading"):
            paragraph = doc.add_heading(block["text"], 0 if kind == "title" else block.get("level", 1))
            paragraph.paragraph_format.keep_with_next = True
            for run in paragraph.runs:
                apply_font(run, 20 if kind == "title" else 14 if block.get("level", 1) == 1 else 11, True)
                run.font.color.rgb = RGBColor.from_string("235675")
        elif kind == "paragraph":
            for line in block["text"].split("\n"):
                paragraph = doc.add_paragraph(line)
                paragraph.paragraph_format.space_after = Pt(6)
                paragraph.paragraph_format.line_spacing = 1.18
                for run in paragraph.runs:
                    apply_font(run, 9 if line.startswith("来源定位：") else 10)
        elif kind == "table":
            headers = block["headers"]
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = "Table Grid"
            table.autofit = False
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                border = OxmlElement("w:" + edge)
                border.set(qn("w:val"), "single")
                border.set(qn("w:sz"), "4")
                border.set(qn("w:color"), "B8C7D1")
                borders.append(border)
            table._tbl.tblPr.append(borders)
            for cell, text in zip(table.rows[0].cells, headers):
                cell.text = text
                shading = OxmlElement("w:shd"); shading.set(qn("w:fill"), "E6EFF4")
                cell._tc.get_or_add_tcPr().append(shading)
            repeat = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for values in block["rows"]:
                for cell, value in zip(table.add_row().cells, values):
                    cell.text = value
            weights = block.get("column_weights", [1] * len(headers))
            for index, column in enumerate(table.columns):
                column.width = Cm(18 * weights[index] / sum(weights))
            for row in table.rows:
                for index, cell in enumerate(row.cells):
                    cell.width = Cm(18 * weights[index] / sum(weights))
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(3)
                        paragraph.paragraph_format.space_before = Pt(3)
                        for run in paragraph.runs:
                            apply_font(run, 8 if len(headers) > 6 else 9)
            if not block["rows"]:
                doc.add_paragraph("无可用完整记录；见数据限制说明。")
            if block.get("note"):
                paragraph = doc.add_paragraph(block["note"])
                for run in paragraph.runs:
                    apply_font(run, 9)
        elif kind == "chart":
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.keep_with_next = True
            paragraph.add_run().add_picture(io.BytesIO(_verify_chart(payload["charts"][block["name"]])), width=Cm(17.6))
            caption = doc.add_paragraph(block["caption"])
            for run in caption.runs:
                apply_font(run, 9)
    paragraph = doc.add_paragraph("冻结报告SHA256：" + payload["frozen_hash"])
    for run in paragraph.runs:
        apply_font(run, 8)
    doc.core_properties.title = payload["mapping"]["报告标题"]
    doc.core_properties.subject = "冻结报告 " + payload["frozen_hash"]
    doc.core_properties.author = "成本智能分析系统"
    doc.core_properties.comments = "任务均为草稿，未发送、未送达。"
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def render_pdf(payload):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle, Image, KeepTogether
    from .model import verify_payload
    verify_payload(payload)
    descriptor = payload["versions"]["renderer"]["font"]
    path = _font_path(descriptor)
    _check_glyphs(payload, path)
    font_name = "CJK-" + descriptor["sha256"][:12]
    with _RENDER_LOCK:
        if font_name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(font_name, path, subfontIndex=descriptor["subfont_index"]))
        styles = {
            "title": ParagraphStyle("ReportTitle", fontName=font_name, fontSize=19, leading=26, textColor=colors.HexColor("#235675"), spaceAfter=14, wordWrap="CJK"),
            "heading": ParagraphStyle("ReportHeading", fontName=font_name, fontSize=13, leading=20, spaceBefore=12, spaceAfter=8, keepWithNext=True, wordWrap="CJK", textColor=colors.HexColor("#235675")),
            "subheading": ParagraphStyle("ReportSubheading", fontName=font_name, fontSize=11, leading=17, spaceBefore=8, spaceAfter=6, keepWithNext=True, wordWrap="CJK"),
            "body": ParagraphStyle("ReportBody", fontName=font_name, fontSize=9.5, leading=15, spaceAfter=7, alignment=TA_LEFT, wordWrap="CJK", splitLongWords=True),
            "cell": ParagraphStyle("ReportCell", fontName=font_name, fontSize=8, leading=12, wordWrap="CJK", splitLongWords=True),
            "small": ParagraphStyle("ReportSmall", fontName=font_name, fontSize=8, leading=12, spaceAfter=5, wordWrap="CJK", splitLongWords=True),
        }
        def paragraph(text, style="body"):
            return Paragraph(escape(str(text)).replace("\n", "<br/>"), styles[style])
        story = []
        width = A4[0] - 3 * cm
        for block in payload["blocks"]:
            kind = block["kind"]
            if kind in ("title", "heading"):
                style = "subheading" if kind == "heading" and block.get("level", 1) > 1 else kind
                story.append(paragraph(block["text"], style))
            elif kind == "paragraph":
                for line in block["text"].split("\n"):
                    style = "small" if line.startswith("来源定位：") else "body"
                    story.append(paragraph(line, style))
            elif kind == "table":
                headers = block["headers"]
                rows = [[paragraph(text, "cell") for text in headers]]
                rows += [[paragraph(text, "cell") for text in row] for row in block["rows"]]
                weights = block.get("column_weights", [1] * len(headers))
                table = LongTable(rows, colWidths=[width * weight / sum(weights) for weight in weights], repeatRows=1, hAlign="LEFT")
                table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E6EFF4")),
                    ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#B8C7D1")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                    ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7FAFC")]),
                ]))
                story.append(table)
                story.append(Spacer(1, 6))
                if not block["rows"]:
                    story.append(paragraph("无可用完整记录；见数据限制说明。", "small"))
                if block.get("note"):
                    story.append(paragraph(block["note"], "small"))
            elif kind == "chart":
                image = Image(io.BytesIO(_verify_chart(payload["charts"][block["name"]])), width=width, height=width * 430 / 1080)
                story.append(KeepTogether([image, paragraph(block["caption"], "small")]))
        story.append(paragraph("冻结报告SHA256：" + payload["frozen_hash"], "small"))
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=1.8 * cm, bottomMargin=1.8 * cm,
                                leftMargin=1.5 * cm, rightMargin=1.5 * cm,
                                title=payload["mapping"]["报告标题"], author="成本智能分析系统",
                                subject="冻结报告 " + payload["frozen_hash"])
        def chrome(canvas, document):
            canvas.saveState()
            canvas.setFont(font_name, 8)
            canvas.setFillColor(colors.HexColor("#557184"))
            canvas.drawString(1.5 * cm, A4[1] - cm, "中药一厂 · 成本分析报告 · 内部")
            review_label = "已审核签发" if payload.get("review_status") == "approved" else "待审核"
            canvas.drawString(1.5 * cm, cm, payload["report_id"] + "  |  " + review_label + "  |  " + payload["frozen_hash"][:12])
            canvas.drawRightString(A4[0] - 1.5 * cm, cm, f"第 {document.page} 页")
            canvas.restoreState()
        doc.build(story, onFirstPage=chrome, onLaterPages=chrome)
        return buffer.getvalue()


def export_report(payload, format="docx"):
    """Render verified frozen input without recomputing numbers or fetching data."""
    if format == "docx":
        return render_docx(payload)
    if format == "pdf":
        return render_pdf(payload)
    raise ValueError("导出格式仅支持docx或pdf")
