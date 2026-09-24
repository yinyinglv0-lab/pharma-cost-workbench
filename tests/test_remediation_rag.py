"""Meaningful acceptance for the additive formal RAG remediation.

The tests use explicit synthetic fixtures, isolated repositories and a deterministic
lexical publication. No repository-original CSV/PDF or local font is required. Real
source/local-BGE acceptance remains separately recorded under artifacts/.../rag and
is never represented as a human-labelled gold set.
"""
from __future__ import annotations

import csv
import hashlib
from io import BytesIO
import threading

import pytest

from enterprise.domain_graph import build_domain_graph
from enterprise.knowledge import KnowledgeAccessError, Repository, preview_file
from enterprise.knowledge_baseline import OUTPUT_FIELDS, SOURCE_FIELDS, build_factory_cost_baseline
from enterprise.knowledge_context import context
from enterprise.knowledge_release import ControlledSearchEngine, ReleaseRepository
from enterprise.security import Principal


ADMIN = Principal('rag-admin', 'RAG管理员', ('knowledge_admin',), ('*',), ('*',))
READER = Principal('rag-reader', 'RAG分析员', ('analyst',), ('中药一厂',), ('银黄口服液',))
OTHER_FACTORY = Principal('rag-other', '二厂分析员', ('analyst',), ('中药二厂',), ('银黄口服液',))


def confirm(repo, text, *, title='受控资料', factory='中药一厂', effective='2026-01-01', products=None,
            visibility='scoped', metadata=None, filename=None):
    products = ['银黄口服液'] if products is None else products
    filename = filename or title + '.txt'
    staged = repo.stage(text.encode('utf-8'), filename, title, products, effective, '工艺', ADMIN.user_id,
                        scope_factories=[factory] if factory else [], visibility=visibility,
                        metadata=metadata or {'authority': 'test-controlled', 'evidence_role': 'document_basis'},
                        principal=ADMIN)
    assert not staged['errors'], staged
    return repo.commit(staged['stage_id'], ADMIN.user_id, '隔离整改测试确认', principal=ADMIN)


def publish(repo, *, dense=False):
    result = ReleaseRepository(repository=repo).publish(principal=ADMIN,
        embedding_model_path='' if not dense else None, require_embeddings=False)
    assert result['status'] == 'published', result
    return result


def test_domain_edges_are_typed_and_have_exact_source_support():
    row = {'chunk_id': 'chunk-1', 'version_id': 'version-1',
           'text': '银黄口服液处方组成：原料名称 金银花；金银花经提取工序进入多功能提取罐，折旧仅作参考。',
           'meta': {'doc_id': 'doc-1', 'filename': '工艺.txt', 'sha256': 'a' * 64,
                    'offset': 20, 'end_offset': 20 + 200, 'scope_products': ['银黄口服液'],
                    'scope_factories': ['中药一厂'], 'effective_from': '2026-01-01',
                    'effective_to': None, 'confirmed_at': '2026-09-21T00:00:00+00:00',
                    'section': '银黄口服液工艺', 'page_spans': []}}
    row['meta']['end_offset'] = row['meta']['offset'] + len(row['text'])
    graph = build_domain_graph([row])
    assert {item['entity_type'] for item in graph.entities.values()} == {'product', 'material', 'process', 'equipment'}
    assert any(item['relation_type'] == 'product_contains_material' and
               item['semantic_status'] == 'semantic_template_supported' for item in graph.relations)
    edge = next(item for item in graph.relations if item['relation_type'] == 'process_uses_equipment')
    support = edge['support']
    assert support['quote'] == row['text'][support['offset'] - row['meta']['offset']:support['end_offset'] - row['meta']['offset']]
    assert support['offset'] >= row['meta']['offset'] and support['end_offset'] <= row['meta']['end_offset']
    assert support['filename'] == '工艺.txt' and support['page'] is None and support['pages'] == []
    assert edge['semantic_status'] == 'semantic_template_supported'


def test_same_release_domain_route_and_controlled_rerank_change_order(tmp_path):
    repo = Repository(tmp_path / 'managed')
    # The first chunk has repeated lexical query terms but no controlled domain
    # entity. The second has the product/material relation that should win after
    # the domain route and explicit rule rank are enabled.
    confirm(repo, '价格上涨 价格上涨 价格上涨 价格上涨 影响 影响 影响', title='泛化价格摘要')
    confirm(repo, '银黄口服液 金银花 提取 工艺', title='银黄工艺依据')
    record = publish(repo)
    engine = ControlledSearchEngine(repository=repo, embedding_model_path='')
    query = '金银花价格上涨对银黄口服液有什么影响'
    baseline, baseline_stats = engine.search(query, principal=READER, product='银黄口服液',
        factory='中药一厂', as_of='2026-06-30', top_k=2, use_domain_graph=False, use_rerank=False)
    formal, formal_stats = engine.search(query, principal=READER, product='银黄口服液',
        factory='中药一厂', as_of='2026-06-30', top_k=2)
    assert record['manifest']['retrieval_schema_version'] == 3
    assert record['manifest']['domain_graph']['status'] == 'ready'
    assert formal_stats['route_generations']['domain_graph'] == record['release_id']
    assert formal_stats['domain_graph_n'] > 0 and formal_stats['reranked']
    trace = formal_stats['rerank_trace']
    assert trace['enabled'] and trace['algorithm'] == 'controlled-domain-lexical-reranker'
    assert trace['neural'] is False and 'changed_from_legacy' in trace
    assert formal[0]['rerank_score'] is not None
    assert formal[0]['route_scores'].get('domain_graph') is not None
    assert baseline_stats['reranked'] is False and baseline
    # Exercise the actual ranker with an adversarial candidate order: the
    # domain relation must change the order, rather than merely add a field.
    version_ids = [item['version_id'] for item in repo.list_documents(principal=ADMIN)]
    rows, graph = engine._load_authorized(record, version_ids)
    ids_by_title = {row['meta']['title']: row['chunk_id'] for row in rows}
    a_id, b_id = ids_by_title['泛化价格摘要'], ids_by_title['银黄工艺依据']
    ranked, direct_trace = engine._controlled_rerank(query, [(a_id, 1.0), (b_id, 0.9)], rows, graph,
                                                      legacy_order=[a_id, b_id])
    assert direct_trace['changed'] and direct_trace['before'] != direct_trace['after']
    assert ranked[0][0] == b_id


def test_domain_graph_respects_effective_date_scope_and_revocation(tmp_path):
    repo = Repository(tmp_path / 'managed')
    old = confirm(repo, '银黄口服液 金银花 提取 工艺', title='晚生效资料', effective='2026-07-01')
    confirm(repo, '银黄口服液 金银花 提取 工艺', title='二厂资料', factory='中药二厂')
    publish(repo)
    engine = ControlledSearchEngine(repository=repo, embedding_model_path='')
    before, before_stats = engine.search('金银花', principal=READER, product='银黄口服液',
        factory='中药一厂', as_of='2026-06-30')
    assert before == [] and before_stats['reason'] == 'no_authorized_effective_published_versions'
    after, after_stats = engine.search('金银花', principal=READER, product='银黄口服液',
        factory='中药一厂', as_of='2026-07-01')
    assert after and all('中药二厂' not in row['meta']['scope_factories'] for row in after)
    assert after_stats['route_generations']['domain_graph'] == after_stats['release_id']
    ReleaseRepository(repository=repo).revoke_document(old['doc_id'], principal=ADMIN, reason='隔离撤回验证')
    revoked, revoked_stats = engine.search('金银花', principal=READER, product='银黄口服液',
        factory='中药一厂', as_of='2026-07-01')
    assert revoked == [] and revoked_stats['authorized_version_ids'] == []
    with pytest.raises(KnowledgeAccessError):
        engine.search('金银花', principal=OTHER_FACTORY, product='银黄口服液',
                      factory='中药一厂', as_of='2026-07-01')


def test_ready_gate_has_bounded_wait_and_does_not_fallback(monkeypatch, tmp_path):
    repo = Repository(tmp_path / 'managed')
    confirm(repo, '银黄口服液 金银花 提取', title='ready')
    publish(repo)
    engine = ControlledSearchEngine(repository=repo, embedding_model_path='')
    monkeypatch.setattr(engine, 'readiness', lambda **kwargs: {
        'ready': False, 'state': 'loading', 'release_id': 'release', 'domain_graph': 'ready',
        'embedding': 'ready', 'wait_timed_out': True})
    rows, stats = engine.search('金银花', principal=READER, product='银黄口服液',
        factory='中药一厂', as_of='2026-06-30', require_ready=True, ready_timeout=0.01)
    assert rows == [] and stats['reason'] == 'retrieval_not_ready'
    assert stats['readiness']['wait_timed_out'] is True


def test_warmup_timeout_returns_while_one_shared_loader_remains_active(monkeypatch, tmp_path):
    import time
    from types import SimpleNamespace
    import enterprise.knowledge_release as kr
    engine = ControlledSearchEngine(root=tmp_path, embedding_model_path='test-local-model')
    ready = threading.Event()
    calls = []
    handle = SimpleNamespace(ready=ready, state='ready', error='')
    monkeypatch.setattr(kr, '_local_model_path', lambda path: tmp_path)
    monkeypatch.setattr(kr, '_model_fingerprint', lambda path: {'test': 'same-model'})
    monkeypatch.setattr(kr, '_model_handle', lambda *args: calls.append(args) or handle)
    started = time.monotonic()
    try:
        state = engine.warmup(principal=ADMIN, wait=True, timeout=0.01)
        assert time.monotonic() - started < 0.5
        assert state['state'] == 'loading' and state['wait_timed_out']
        engine.warmup(principal=ADMIN, wait=False)
        assert len(calls) == 1
    finally:
        ready.set()
    assert engine.warmup(principal=ADMIN, wait=True, timeout=1)['state'] == 'ready'


def test_legacy_release_can_be_ready_without_domain_schema(monkeypatch, tmp_path):
    engine = ControlledSearchEngine(root=tmp_path, embedding_model_path='')
    monkeypatch.setattr(engine.releases, '_raw_release', lambda *args: {
        'release_id': 'old-release', 'manifest': {'schema_version': 2, 'embedding': {'status': 'disabled'}}})
    ready = engine.readiness(principal=READER)
    assert ready['ready'] and ready['domain_graph'] == 'disabled_legacy'
    assert ready['state'] == 'ready_lexical'


def test_context_keeps_source_location_and_real_route_scores(tmp_path):
    repo = Repository(tmp_path / 'managed')
    confirm(repo, '银黄口服液 金银花 提取', title='source')
    publish(repo)
    sources, stats = context('银黄口服液', '2026-06', '金银花', repository=repo,
                             principal=READER, factory='中药一厂', return_stats=True)
    assert sources and stats['rerank_trace']
    source = sources[0]
    assert source['text'] and source['source']['offset'] is not None
    assert source['source']['end_offset'] > source['source']['offset']
    assert source['source']['page'] is None and source['source']['pages'] == []
    assert set(('rrf_score', 'rerank_score', 'vector_score', 'bm25_score', 'graph_score')) <= set(source['retrieval_scores'])
    assert source['retrieval_scores']['rrf_score'] is not None


def test_pdf_docx_txt_parser_range_is_explicit():
    # Build the parser inputs in memory so source-bundle execution never reaches
    # excluded repository originals. This still exercises PDF/DOCX/TXT parsing,
    # ordering, encoding, and page-range metadata.
    from reportlab.pdfgen import canvas
    pdf_buffer = BytesIO()
    pdf_document = canvas.Canvas(pdf_buffer)
    pdf_document.drawString(30, 700, 'Synthetic process evidence')
    pdf_document.showPage()
    pdf_document.save()
    pdf = preview_file(pdf_buffer.getvalue(), 'synthetic-process.pdf')
    assert not pdf['errors'] and pdf['parser'] == 'pypdf' and pdf['metadata']['page_count'] == 1
    assert 'Synthetic process evidence' in pdf['text']
    from docx import Document
    doc = Document()
    doc.add_paragraph('银黄口服液 金银花 提取工艺')
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = '设备'
    table.rows[0].cells[1].text = '多功能提取罐'
    merged = table.add_row().cells
    merged[0].merge(merged[1]).text = '合并说明：按原表保留，页码不推测'
    doc.add_paragraph('表后核查记录')
    buffer = BytesIO()
    doc.save(buffer)
    docx = preview_file(buffer.getvalue(), '受控工艺.docx')
    assert not docx['errors'] and docx['parser'] == 'python-docx' and docx['metadata']['table_count'] == 1
    assert docx['text'].index('金银花') < docx['text'].index('多功能提取罐') < docx['text'].index('表后核查记录')
    assert '合并说明' in docx['text'] and 'page_count' not in docx['metadata']
    txt_utf8 = preview_file('银黄口服液\n金银花'.encode('utf-8'), '受控.txt')
    txt_gbk = preview_file('银黄口服液\n金银花'.encode('gb18030'), '受控-gbk.txt')
    assert not txt_utf8['errors'] and txt_utf8['metadata']['encoding'] == 'utf-8-sig'
    assert not txt_gbk['errors'] and txt_gbk['metadata']['encoding'] == 'gb18030'
    for parsed in (pdf, docx, txt_utf8, txt_gbk):
        assert parsed['text'] and parsed['text_sha256']


def test_synthetic_semantic_recipe_flow_equipment_chain_survives_section_chunking(tmp_path):
    from enterprise.knowledge_release import _chunks
    repo = Repository(tmp_path / 'managed')
    rows = []
    versions = {}
    fixtures = [
        ('synthetic_recipe_a.txt', '合成银黄口服液配方', ['银黄口服液'],
         '一、银黄口服液配方\n合成说明，仅供测试。\n二、处方组成\n原料名称 金银花 10g\n'),
        ('synthetic_process.txt', 'synthetic process equipment guide',
         ['银黄口服液', '板蓝根颗粒', '六味地黄胶囊'],
         '银黄口服液生产工艺\n金银花→提取\n提取使用设备：多功能提取罐\n'
         '板蓝根颗粒生产工艺\n糊精→混合\n混合使用设备：湿法制粒机\n'
         '粉碎 粉碎收率 >= 95%\n灭菌 灭菌合格率 >= 99%\n'
         '六味地黄胶囊生产工艺\n熟地黄→干燥\n'),
        ('synthetic_recipe_b.txt', '合成板蓝根颗粒配方', ['板蓝根颗粒'],
         '一、板蓝根颗粒配方\n合成说明，仅供测试。\n二、处方组成\n原料名称 板蓝根 20g\n原料名称 糊精 10g\n'),
    ]
    for filename, title, products, text in fixtures:
        staged = repo.stage(text.encode('utf-8'), filename, title, products, '2025-01-01',
                            '产品配方' if 'recipe' in filename else '生产工艺', ADMIN.user_id,
                            scope_factories=['中药一厂'], visibility='scoped', principal=ADMIN)
        assert not staged['errors']
        version = repo.commit(staged['stage_id'], ADMIN.user_id, '隔离合成原文测试', principal=ADMIN)
        versions[version['version_id']] = version
        rows.extend(_chunks(version, {'name': 'product-section-character-window', 'version': 2,
                                       'size': 900, 'overlap': 150}))
    graph = build_domain_graph(rows)
    semantic = [edge for edge in graph.relations if edge['semantic_status'] == 'semantic_template_supported']
    triples = {(graph.entities[e['source_entity_id']]['name'], e['relation_type'],
                graph.entities[e['target_entity_id']]['name']) for e in semantic}
    assert ('银黄口服液', 'product_contains_material', '金银花') in triples
    assert ('金银花', 'material_undergoes_process', '提取') in triples
    # A synthetic cross-source recipe -> material -> process -> equipment path.
    assert ('板蓝根颗粒', 'product_contains_material', '糊精') in triples
    assert ('糊精', 'material_undergoes_process', '混合') in triples
    assert ('混合', 'process_uses_equipment', '湿法制粒机') in triples
    assert ('粉碎', 'process_has_controlled_metric', '粉碎收率') in triples
    assert ('灭菌', 'process_has_controlled_metric', '粉碎收率') not in triples
    assert any(context_source['chunk_id'] != edge['support']['chunk_id']
               for edge in semantic for context_source in edge['support'].get('context_sources', []))
    for edge in semantic:
        support = edge['support']
        original = versions[support['version_id']]['text']
        assert original[support['offset']:support['end_offset']] == support['quote']
        assert support['scope_factories'] == ['中药一厂']
        for context_source in support.get('context_sources', []):
            assert original[context_source['offset']:context_source['end_offset']] == context_source['quote']


def test_general_and_nonadjacent_mentions_never_become_semantic_predicates():
    text = ('药品生产质量管理规范：金银花、混合、湿法制粒机须登记。' + '通用要求。' * 80
            + '银黄口服液仅为索引词。配方无关。')
    graph = build_domain_graph([{'chunk_id': 'gmp', 'version_id': 'v', 'text': text,
        'meta': {'doc_id': 'doc', 'filename': 'gmp.txt', 'offset': 0, 'end_offset': len(text),
                 'title': 'GMP通用说明', 'scope_products': ['银黄口服液'], 'scope_factories': ['中药一厂']}}])
    assert graph.relations
    assert all(edge['semantic_status'] == 'co_occurrence_only' for edge in graph.relations)
    assert graph.domain_search('金银花')[1]['relation_hits'] == {}


def test_keyword_synonym_has_measurable_bm25_effect(tmp_path):
    repo = Repository(tmp_path / 'managed')
    confirm(repo, '灌装 (10ml/支, 灌封一体机)', title='设备')
    publish(repo)
    engine = ControlledSearchEngine(repository=repo, embedding_model_path='')
    rows, stats = engine.search('灌装封口机', principal=READER, product='银黄口服液',
                               factory='中药一厂', as_of='2026-06-30')
    assert rows and stats['keyword_expansions']
    raw, graph = engine._load_authorized(engine.releases._raw_release(), set(stats['authorized_version_ids']))
    before = dict(engine._bm25('灌装封口机', raw))
    assert rows[0]['route_scores']['bm25'] > before.get(rows[0]['chunk_id'], 0)


def test_empty_scan_pdf_and_doc_are_rejected_without_fake_page_numbers():
    from pypdf import PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    raw = BytesIO()
    writer.write(raw)
    assert preview_file(raw.getvalue(), '扫描.pdf')['errors']
    assert preview_file(b'legacy doc bytes', '旧文档.doc')['errors']


@pytest.fixture
def synthetic_cost_sources(tmp_path):
    root = tmp_path / 'synthetic-cost-inputs'
    root.mkdir()
    sources = {}
    for marker, encoding in (('a', 'utf-8-sig'), ('b', 'gb18030')):
        filename = f'synthetic_cost_{marker}.csv'
        rows = [dict(zip(SOURCE_FIELDS, (
            f'synthetic-factory-{marker}', f'synthetic-product-{index // 6 + 1}', 'synthetic-spec',
            f'2026-{index % 6 + 1:02d}', '100', '1.2500', '0.500', '0.25', '2.0000', '200.0000',
        ))) for index in range(18)]
        with (root / filename).open('w', encoding=encoding, newline='') as stream:
            writer = csv.DictWriter(stream, fieldnames=SOURCE_FIELDS, lineterminator='\n')
            writer.writeheader()
            writer.writerows(rows)
        sources[filename] = {'rows': rows, 'bytes': (root / filename).read_bytes()}
    return root, sources


def test_derived_baseline_preserves_source_hashes_and_gaps(tmp_path, synthetic_cost_sources):
    root, sources = synthetic_cost_sources
    output = tmp_path / 'baseline.csv'
    result = build_factory_cost_baseline(root, output, filenames=tuple(sources))
    metadata = result['metadata']
    assert metadata['row_count'] == 36
    assert {item['factory'][0] for item in metadata['source_files']} == {
        'synthetic-factory-a', 'synthetic-factory-b'}
    assert {item['filename'] for item in metadata['source_files']} == set(sources)
    for item in metadata['source_files']:
        raw = sources[item['filename']]['bytes']
        assert item['row_count'] == 18 and item['bytes'] == len(raw)
        assert item['sha256'] == hashlib.sha256(raw).hexdigest()
        assert (root / item['filename']).read_bytes() == raw
    with output.open(encoding='utf-8', newline='') as stream:
        reader = csv.DictReader(stream)
        assert tuple(reader.fieldnames) == OUTPUT_FIELDS
        derived_rows = list(reader)
    assert len(derived_rows) == 36
    expected_rows = {(name, str(index)): row for name, source in sources.items()
                     for index, row in enumerate(source['rows'], start=2)}
    assert {(row['source_file'], row['source_row']) for row in derived_rows} == set(expected_rows)
    for row in derived_rows:
        assert {field: row[field] for field in SOURCE_FIELDS} == expected_rows[(row['source_file'], row['source_row'])]
        assert row['source_sha256'] == hashlib.sha256(sources[row['source_file']]['bytes']).hexdigest()
        assert row['derivation_status'] == 'observed_cost_summary'
    assert metadata['output_sha256'] == hashlib.sha256(output.read_bytes()).hexdigest()
    assert metadata['derivation']['formula'] is None
    assert '模拟' in metadata['known_gaps'][0]
    assert '采购价' in metadata['claim_boundary'] and '收率' in metadata['claim_boundary']
