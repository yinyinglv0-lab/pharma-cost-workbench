"""Isolated version, duplicate, temporal, parser and confirmation acceptance."""
from datetime import datetime, timezone
import hashlib
import io
import sqlite3
import zipfile

import pytest

from enterprise.knowledge import Repository, KnowledgeError, VersionConflict, preview_file, MAX_BYTES


@pytest.fixture
def repo(tmp_path):
    return Repository(tmp_path/'managed')


def stage(repo, text='第一条制度\n第二条制度', **kw):
    args=dict(content=text.encode('utf-8'),filename='制度.txt',title='制度',scope_products=[],
              effective_from='2026-01-01',category='通用制度',actor='tester')
    args.update(kw)
    return repo.stage(**args)


def commit(repo, text='第一条制度\n第二条制度', **kw):
    pending=stage(repo,text,**kw)
    assert not pending['errors'],pending
    return repo.commit(pending['stage_id'],'reviewer','核对原文后确认')


def test_read_only_empty_and_stage_is_not_effective(repo):
    assert not repo.root.exists()
    assert repo.list_documents()==[] and repo.events()==[]
    assert repo.effective_versions('银黄口服液','2026-05-01')==[]
    assert not repo.root.exists()
    pending=stage(repo)
    assert pending['base_version']==0 and pending['change']['kind']=='new_document'
    assert repo.list_documents()==[] and not repo.blob_dir.exists()
    confirmed=repo.commit(pending['stage_id'],'user','确认')
    assert confirmed['version']==1 and confirmed['index_status']=='catalog_only'
    assert len(confirmed['sha256'])==64
    assert repo.read_blob(confirmed['sha256']).decode('utf-8')=='第一条制度\n第二条制度'
    assert repo.commit(pending['stage_id'],'user','重复确认')['version_id']==confirmed['version_id']
    assert len(repo.history(confirmed['doc_id']))==1


def test_immutable_update_history_blob_and_diff(repo):
    first=commit(repo)
    pending=stage(repo,'第一条制度\n修订后的第二条',doc_id=first['doc_id'],effective_from='2026-05-01')
    assert pending['base_version']==1 and pending['change']['kind']=='new_version'
    assert '-第二条制度' in pending['diff'] and '+修订后的第二条' in pending['diff']
    second=repo.commit(pending['stage_id'],'another','修订条款')
    assert second['version']==2
    assert repo.get(first['doc_id'],1)['text']==first['text']
    assert repo.get(version_id=first['version_id'])['sha256']==first['sha256']
    assert repo.read_blob(first['sha256'])!=repo.read_blob(second['sha256'])
    assert len(repo.history(first['doc_id']))==2
    with sqlite3.connect(repo.db_path) as conn:
        with pytest.raises(sqlite3.IntegrityError):conn.execute('UPDATE versions SET text=?',('tampered',))
        with pytest.raises(sqlite3.IntegrityError):conn.execute('DELETE FROM versions')


def test_exact_duplicate_skips_parser_and_all_writes(repo,monkeypatch):
    first=commit(repo)
    before=repo.db_path.read_bytes()
    events=repo.events()
    def fail(*args):raise AssertionError('duplicate parser must not run')
    monkeypatch.setattr('enterprise.knowledge.preview_file',fail)
    pending=stage(repo,doc_id=first['doc_id'])
    assert pending['stage_id'] is None
    assert pending['change']['kind']=='binary_duplicate'
    assert repo.db_path.read_bytes()==before and repo.events()==events
    assert len(list(repo.blob_dir.iterdir()))==1


def test_same_parsed_text_different_bytes_same_business_version_skips(repo):
    first=commit(repo,'甲\n乙')
    before=len(repo.events())
    pending=stage(repo,content='甲\r\n乙\r\n'.encode('gb18030'),filename='另一份.txt',doc_id=first['doc_id'])
    assert pending['change']['kind']=='text_duplicate'
    assert '不必更新' in pending['change']['summary']
    assert pending['change']['duplicate_of']['version_id']==first['version_id']
    assert len(repo.events())==before and len(repo.list_documents())==1


def test_optimistic_base_version_conflict(repo):
    first=commit(repo)
    a=stage(repo,'修订A',doc_id=first['doc_id'])
    b=stage(repo,'修订B',doc_id=first['doc_id'])
    repo.commit(a['stage_id'],'A','先确认')
    with pytest.raises(VersionConflict,match='新版本'):
        repo.commit(b['stage_id'],'B','后确认')
    assert len(repo.history(first['doc_id']))==2


def test_concurrent_new_document_same_title_conflict(repo):
    a=stage(repo,'A')
    b=stage(repo,'B')
    repo.commit(a['stage_id'],'A','确认')
    with pytest.raises(VersionConflict,match='同标题'):
        repo.commit(b['stage_id'],'B','确认')


def test_concurrent_same_content_distinct_documents_share_blob(repo):
    a=stage(repo,'共同条款',title='甲')
    b=stage(repo,'共同条款',title='乙')
    first=repo.commit(a['stage_id'],'A','确认')
    second=repo.commit(b['stage_id'],'B','确认')
    assert second['status']=='confirmed'
    assert first['doc_id']!=second['doc_id'] and first['sha256']==second['sha256']
    assert len(repo.list_documents())==2
    assert len(list(repo.blob_dir.iterdir()))==1


def test_same_body_metadata_revision_preserves_identity_and_blob(repo):
    first=commit(repo,scope_products=['A'])
    pending=stage(repo,doc_id=first['doc_id'],scope_products=['A','B'],effective_from='2026-02-01',metadata={'owner':'质量部'})
    assert pending['change']['kind']=='metadata_revision' and pending['diff']==''
    assert pending['change']['metadata_changes']['scope_products']=={'before':['A'],'after':['A','B']}
    second=repo.commit(pending['stage_id'],'reviewer','范围获批')
    assert second['version']==2 and second['doc_id']==first['doc_id']
    assert second['text']==first['text'] and second['sha256']==first['sha256']
    assert repo.get(version_id=first['version_id'])['scope_products']==['A']
    assert repo.effective_versions('B','2026-02-01')[0]['version_id']==second['version_id']
    assert len(list(repo.blob_dir.iterdir()))==1


def test_stage_metadata_tampering_rejected(repo):
    pending=stage(repo)
    with sqlite3.connect(repo.db_path) as conn:
        conn.execute("UPDATE stages SET payload=replace(payload, '通用制度', '篡改类别') WHERE stage_id=?",(pending['stage_id'],))
    with pytest.raises(KnowledgeError,match='metadata校验失败'):
        repo.commit(pending['stage_id'],'reviewer','确认')
    assert repo.history()==[]


def test_future_version_cannot_pollute_past_scope_and_overlap(repo):
    first=commit(repo,'初版')
    future=commit(repo,'新版',doc_id=first['doc_id'],effective_from='2026-07-01',scope_products=['银黄口服液'])
    assert repo.effective_versions('板蓝根颗粒','2026-06-30')[0]['version_id']==first['version_id']
    assert repo.effective_versions('银黄口服液','2026-07-01')[0]['version_id']==future['version_id']
    # Latest effective version changes scope; do not resurrect the older generic version for other products.
    assert repo.effective_versions('板蓝根颗粒','2026-07-01')==[]
    assert repo.effective_versions('银黄口服液','2026-07-01',known_at=first['confirmed_at'])[0]['version_id']==first['version_id']


def test_effective_order_is_business_date_not_commit_order(repo):
    first=commit(repo,'五月',effective_from='2026-05-01')
    later_confirmed=commit(repo,'三月追补',doc_id=first['doc_id'],effective_from='2026-03-01')
    assert repo.effective_versions('任意','2026-04-01')[0]['version_id']==later_confirmed['version_id']
    assert repo.effective_versions('任意','2026-06-01')[0]['version_id']==first['version_id']


def test_date_bounds_and_operation_filters(repo):
    first=commit(repo,'限期',effective_from='2026-03-01',effective_to='2026-03-31')
    assert len(repo.effective_versions('产品','2026-03-31'))==1
    assert repo.effective_versions('产品','2026-04-01')==[]
    assert len(repo.history(effective_from='2026-03-01',effective_to='2026-03-31'))==1
    assert repo.history(effective_from='2026-04-01')==[]
    op=first['confirmed_at'][:10]
    assert len(repo.list_documents(operation_from=op,operation_to=op))==1
    assert len(repo.events(operation_from=op,operation_to=op,effective_from='2026-03-01'))==2
    assert repo.events(operation_to='2000-01-01')==[]


@pytest.mark.parametrize('kwargs,reason',[
    ({'filename':'x.doc'},'.docx'),({'filename':'x.exe'},'仅支持'),
    ({'content':b''},'不能为空'),({'content':b'x'*(MAX_BYTES+1)},'20MB'),
    ({'effective_from':'2026-02-30'},'有效日期'),({'effective_from':'2026-1-1'},'YYYY-MM-DD'),
    ({'effective_to':'2025-12-31'},'不得早于'),({'scope_products':'银黄口服液'},'列表'),
    ({'actor':''},'必填'),({'doc_id':'missing'},'不存在'),
])
def test_bad_input_never_creates_stage(repo,kwargs,reason):
    result=stage(repo,**kwargs)
    assert reason in ' '.join(result['errors'])
    assert result['stage_id'] is None and not repo.db_path.exists()
    assert not repo.blob_dir.exists()  # Only the shared maintenance lock may exist.


def test_filename_and_blob_path_safety(repo):
    result=commit(repo,filename='C:\\temp\\..\\..\\制度.txt')
    assert result['filename']=='制度.txt'
    assert (repo.blob_dir/result['sha256']).is_file()
    for path in ('../../secret','A'*64,'abc'):
        with pytest.raises(KnowledgeError):repo.read_blob(path)
    (repo.blob_dir/result['sha256']).write_bytes(b'corrupted')
    with pytest.raises(KnowledgeError,match='哈希校验失败'):repo.read_blob(result['sha256'])


def test_csv_original_lines_and_gb18030():
    content='材料,说明\r\n金银花,"一行\n二行"\r\n'.encode('gb18030')
    result=preview_file(content,'材料.csv')
    assert not result['errors']
    assert result['metadata']['row_count']==2
    assert result['metadata']['preview_rows'][1][1]=='一行\n二行'
    assert '\n金银花,' in result['text']
    assert result['metadata']['encoding']=='gb18030'


def test_docx_paragraph_table_order_and_binary_change(repo):
    from docx import Document
    doc=Document();doc.add_paragraph('表格前')
    table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='药材';table.cell(0,1).text='金银花'
    doc.add_paragraph('表格后')
    a=io.BytesIO();doc.save(a)
    parsed=preview_file(a.getvalue(),'规则.docx')
    assert not parsed['errors']
    assert parsed['text']=='表格前\n药材\t金银花\n表格后'
    p=stage(repo,content=a.getvalue(),filename='规则.docx')
    repo.commit(p['stage_id'],'u','确认')
    doc.core_properties.title='只改文档属性'
    b=io.BytesIO();doc.save(b)
    p=stage(repo,content=b.getvalue(),filename='规则.docx')
    assert p['change']['kind']=='text_duplicate'


def test_local_pdf_parser():
    from reportlab.pdfgen import canvas
    content = io.BytesIO()
    document = canvas.Canvas(content)
    document.drawString(30, 700, 'Local knowledge evidence')
    document.showPage()
    document.save()
    result = preview_file(content.getvalue(), 'original.pdf')
    assert not result['errors'], result
    assert result['parser'] == 'pypdf'
    assert 'Local knowledge evidence' in result['text'] and '[第1页]' in result['text']
    assert result['metadata']['page_count'] == 1


def test_empty_scan_pdf_is_rejected():
    from reportlab.pdfgen import canvas
    content = io.BytesIO()
    document = canvas.Canvas(content)
    document.showPage()
    document.save()
    result = preview_file(content.getvalue(), 'scan.pdf')
    assert result['parser'] == 'pypdf'
    assert 'OCR' in ' '.join(result['errors'])


def test_confirm_reason_and_unknown_stage(repo):
    pending=stage(repo)
    with pytest.raises(KnowledgeError,match='原因必填'):repo.commit(pending['stage_id'],'u','')
    with pytest.raises(KnowledgeError,match='预览不存在'):repo.commit('unknown','u','原因')
    assert repo.list_documents()==[]


def test_streamlit_page_readonly_and_explicit_confirmation(repo, monkeypatch):
    from pathlib import Path
    from streamlit.testing.v1 import AppTest
    import app_pages._shared as shared
    from enterprise.application import Application
    from enterprise.knowledge_release import ReleaseRepository
    from enterprise.security import Principal

    principal = Principal('knowledge-ui-test', '知识管理员', ('knowledge_admin',), ('*',), ('*',))
    service = Application(principal, repo.root)
    monkeypatch.setattr(shared, 'streamlit_principal', lambda: principal)
    monkeypatch.setattr(shared, 'Application', lambda actor: service)
    page = Path(__file__).resolve().parents[1] / 'app_pages' / 'knowledge.py'
    script = ('import streamlit as st\n'
              f"st.navigation([st.Page({str(page)!r}, title='知识文档与版本')]).run()\n")
    app = AppTest.from_string(script, default_timeout=30).run()
    releases = ReleaseRepository(repository=repo)
    assert not app.exception and not app.error
    assert repo.history() == [], 'opening the page must not register or confirm originals'
    assert releases.history(principal=principal) == [], 'opening the page must not publish an index'
    assert not repo.blob_dir.exists()

    pending = stage(repo, visibility='public', principal=principal)
    app.session_state['knowledge_stage'] = pending
    app.run()
    assert not app.exception and repo.history() == []
    next(x for x in app.text_input if x.label == '登记 / 更新原因').set_value('页面确认测试')
    next(x for x in app.button if x.label == '确认登记此版本').click().run()
    assert not app.exception and app.error and repo.history() == []
    next(x for x in app.checkbox if x.label == '已核对正文、适用范围、生效日期与差异').check()
    next(x for x in app.button if x.label == '确认登记此版本').click().run()
    assert not app.exception and not app.error
    versions = repo.history()
    assert len(versions) == 1 and versions[0]['confirmed_by'] == principal.user_id
    assert 'knowledge_stage' not in app.session_state
    assert releases.history(principal=principal) == [], 'confirmation alone must not publish an index'
