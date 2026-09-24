"""Read-only knowledge UI regressions; repositories live only under tmp_path."""
from __future__ import annotations

import ast
from dataclasses import replace
import io
from pathlib import Path
import sys
from types import SimpleNamespace

import pytest
import streamlit as st
from streamlit.testing.v1 import AppTest

from enterprise.application import Application
from enterprise.knowledge import Repository
from enterprise.security import Principal
from enterprise.tabular_knowledge import KNOWLEDGE_TYPES, MARKET_HEADERS


ROOT = Path(__file__).resolve().parents[1]
PAGE = ROOT / 'app_pages' / 'knowledge.py'
ADMIN = Principal('browser-admin', '测试管理员', ('knowledge_admin',), ('*',), ('*',))
READER = replace(ADMIN, user_id='browser-reader', roles=('analyst',),
                 factories=('一厂',), products=('甲产品',))


@pytest.fixture(scope='module')
def helpers():
    """Load pure/page-render helpers without executing the Streamlit page body."""
    tree = ast.parse(PAGE.read_text(encoding='utf-8'), filename=str(PAGE))
    body = []
    for node in tree.body:
        if isinstance(node, ast.Assign) and any(isinstance(target, ast.Tuple) for target in node.targets):
            break  # principal, app = page_context(...) is the page entry point.
        body.append(node)
    namespace = {}
    exec(compile(ast.Module(body=body, type_ignores=[]), str(PAGE), 'exec'), namespace)
    return SimpleNamespace(**namespace)


@pytest.fixture
def repo(tmp_path):
    return Repository(tmp_path / 'managed', principal=ADMIN)


def confirm(repo, title='可见正文', text='甲产品 提取工艺 正文', filename='visible.txt', category='生产工艺',
            *, product='甲产品', factory='一厂', content=None, metadata=None, doc_id=None):
    pending = repo.stage(content if content is not None else text.encode('utf-8'), filename, title,
                         [product], '2026-01-01', category, ADMIN.user_id, doc_id=doc_id,
                         scope_factories=[factory], visibility='scoped', metadata=metadata or {})
    assert not pending['errors'], pending
    return repo.commit(pending['stage_id'], ADMIN.user_id, '临时测试资料确认')


def open_page(monkeypatch, repo, principal=READER):
    import app_pages._shared as shared
    monkeypatch.setattr(shared, 'streamlit_principal', lambda: principal)
    monkeypatch.setattr(shared, 'Application', lambda actor: Application(actor, repo.root))
    # Use a real navigation entry point but no service, cloud, or original DB.
    script = 'import streamlit as st\n' + f'st.navigation([st.Page({str(PAGE)!r})]).run()\n'
    app = AppTest.from_string(script, default_timeout=30).run()
    assert not app.exception
    return app


def literal_content(app):
    return '\n'.join(str(item.value) for name in ('text', 'text_area', 'caption', 'info', 'warning', 'error')
                     for item in getattr(app, name))


def test_csv_records_preserve_every_cell_record_and_duplicate_header(helpers):
    source = 'id,id,\n001,"line one\nline two",=SUM(A1)\n\n003,<script>alert(1)</script>\n004,x,y,extra'
    rows, ragged = helpers._csv_records(source)
    assert len(rows) == 5
    assert rows[0] == {'CSV记录': 1, '字段数': 3, '列 1': 'id', '列 2': 'id', '列 3': '', '列 4': None}
    assert rows[1]['列 1'] == '001' and rows[1]['列 2'] == 'line one\nline two'
    assert rows[1]['列 3'] == '=SUM(A1)'
    assert rows[2]['字段数'] == 0 and rows[2]['列 1'] is None
    assert rows[3]['列 2'] == '<script>alert(1)</script>'
    assert rows[4]['列 4'] == 'extra'
    assert ragged == [3, 4, 5]


@pytest.mark.parametrize('text', ['a,b\n1,"unterminated', 'a,b\n1,"x"oops'])
def test_malformed_csv_rejects_whole_table_instead_of_skipping_rows(helpers, text):
    with pytest.raises(ValueError, match='未展示部分结果'):
        helpers._csv_records(text)


def test_filters_use_only_authorized_documents_and_handle_empty_combination(repo, monkeypatch):
    csv = confirm(repo, title='授权表格', filename='visible.csv', text='编号,数值\n001,2', category='行业基准')
    txt = confirm(repo, title='授权工艺')
    confirm(repo, title='不应泄露的外厂标题', filename='secret.txt', category='外厂秘密类别', factory='二厂')
    before = repo.history()
    app = open_page(monkeypatch, repo)
    # 类别选项=统一目录（含暂无文档的类别）+文档数标注；不得泄露未授权文档的类别
    category_options = app.selectbox(key='knowledge_browse_category').options
    assert '全部授权类别' in category_options
    assert any('生产工艺（1 份）' in option for option in category_options)
    assert any('行业基准（1 份）' in option for option in category_options)
    assert any('异常处理记录' in option for option in category_options)
    assert not any('外厂秘密类别' in option for option in category_options)
    assert app.selectbox(key='knowledge_browse_format').options == ['全部授权格式', 'PDF', 'Word(DOCX)', 'TXT', 'CSV']
    assert set(app.dataframe[0].value['文档']) == {'授权表格', '授权工艺'}
    assert '不应泄露' not in literal_content(app)
    assert not app.file_uploader
    app.selectbox(key='knowledge_browse_category').set_value('行业基准').run()
    assert not app.exception and list(app.dataframe[0].value['文档']) == ['授权表格']
    assert app.selectbox(key='knowledge_browse_doc').value == csv['doc_id']
    app.selectbox(key='knowledge_browse_format').set_value('txt').run()
    assert not app.exception and not app.error
    assert not any(widget.key == 'knowledge_browse_doc' for widget in app.selectbox)
    assert any('没有文档' in item.value for item in app.info)
    app.selectbox(key='knowledge_browse_category').set_value(None).run()
    assert not app.exception and app.selectbox(key='knowledge_browse_doc').value == txt['doc_id']
    assert repo.history() == before


def test_csv_readable_display_raw_download_version_hash_and_no_write(repo, monkeypatch):
    source = '代码,备注,备注\n001,"多行\n字段",=1+1\n\n002,缺列\n003,x,y,额外字段'
    selected = confirm(repo, filename='rows.csv', text=source, category='其他')
    captured = []
    original = st.download_button
    def download(label, data, **kwargs):
        captured.append((data, kwargs))
        return original(label, data, **kwargs)
    monkeypatch.setattr(st, 'download_button', download)
    before = repo.history()
    app = open_page(monkeypatch, repo)
    assert not app.error and app.warning
    table = app.dataframe[1].value
    assert len(table) == 5 and table.iloc[1]['列 1'] == '001'
    assert table.iloc[1]['列 2'] == '多行\n字段' and table.iloc[1]['列 3'] == '=1+1'
    assert table.iloc[2]['字段数'] == 0 and table.iloc[4]['列 4'] == '额外字段'
    assert app.text_area(key=f"knowledge_read_{selected['version_id']}").value == selected['text']
    assert captured[0][0] == source.encode('utf-8') and captured[0][1]['file_name'] == 'rows.csv'
    assert selected['sha256'] in literal_content(app)
    assert selected['text_sha256'] in literal_content(app)
    assert selected['version_id'] in literal_content(app)
    assert repo.history() == before


def test_csv_pagination_does_not_omit_later_records(repo, monkeypatch):
    source = 'id,value\n' + '\n'.join(f'{number:03},x' for number in range(205))
    selected = confirm(repo, filename='many.csv', text=source)
    app = open_page(monkeypatch, repo)
    key = f"knowledge_csv_page_{selected['version_id']}"
    assert app.selectbox(key=key).options == ['1', '2', '3']
    assert len(app.dataframe[1].value) == 100
    app.selectbox(key=key).set_value(3).run()
    assert not app.exception
    assert list(app.dataframe[1].value['CSV记录']) == list(range(201, 207))
    assert app.dataframe[1].value.iloc[-1]['列 1'] == '204'


def test_malformed_confirmed_csv_falls_back_without_partial_table(repo, monkeypatch):
    selected = confirm(repo, filename='rows.csv', text='id,value\n001,x')
    original = Repository.get
    damaged = 'id,value\n001,"unterminated'
    def get(self, *args, **kwargs):
        row = original(self, *args, **kwargs)
        return {**row, 'text': damaged} if row and row['version_id'] == selected['version_id'] else row
    monkeypatch.setattr(Repository, 'get', get)  # Simulate a legacy damaged value; never modify SQLite.
    app = open_page(monkeypatch, repo)
    assert any('未展示部分结果' in item.value for item in app.warning)
    assert len(app.dataframe) == 1
    assert app.text_area(key=f"knowledge_read_{selected['version_id']}").value == damaged
    assert len(app.get('download_button')) == 1


def pdf_content():
    # Core PDF dependency only; no extra renderer/reportlab install is required.
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
    writer = PdfWriter()
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                             NameObject('/Subtype'): NameObject('/Type1'),
                             NameObject('/BaseFont'): NameObject('/Helvetica')})
    for value in ('First source page', 'Second source page'):
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):
            DictionaryObject({NameObject('/F1'): font})})
        stream = DecodedStreamObject()
        stream.set_data(f'BT /F1 12 Tf 30 700 Td ({value}) Tj ET'.encode('ascii'))
        page[NameObject('/Contents')] = stream
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def test_pdf_without_optional_component_uses_download_and_page_markers(repo, monkeypatch):
    content = pdf_content()
    selected = confirm(repo, filename='original.pdf', content=content, category='法规原文')
    monkeypatch.setitem(sys.modules, 'streamlit_pdf', None)
    monkeypatch.setattr(st, 'pdf', lambda *args, **kwargs: pytest.fail('missing optional viewer must not be called'))
    app = open_page(monkeypatch, repo)
    assert not app.error and any('未启用 PDF' in item.value for item in app.info)
    assert any('[第1页]' in item.value and '[第2页]' in item.value for item in app.text)
    assert app.text_area(key=f"knowledge_read_{selected['version_id']}").value == selected['text']
    assert len(app.get('download_button')) == 1


@pytest.mark.parametrize('fail', [False, True])
def test_pdf_verified_component_receives_authorized_bytes_and_failure_is_safe(repo, monkeypatch, fail):
    content = pdf_content()
    selected = confirm(repo, filename='original.pdf', content=content)
    monkeypatch.setitem(sys.modules, 'streamlit_pdf', SimpleNamespace(pdf_viewer=lambda **kwargs: None))
    captured = []
    def pdf(data, **kwargs):
        captured.append((data, kwargs))
        if fail:
            raise RuntimeError('component unavailable')
    monkeypatch.setattr(st, 'pdf', pdf)
    app = open_page(monkeypatch, repo)
    assert not app.error and captured == [(content, {'height': 500, 'key': f"knowledge_pdf_{selected['version_id']}"})]
    assert app.text_area(key=f"knowledge_read_{selected['version_id']}").value == selected['text']
    assert bool([item for item in app.info if '暂不可用' in item.value]) is fail
    assert len(app.get('download_button')) == 1


def test_docx_safe_plain_paragraph_table_order_without_html_execution(repo, monkeypatch):
    from docx import Document
    doc = Document()
    doc.add_paragraph('<script>alert("before")</script>')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = '编号'; table.cell(0, 1).text = '资料'
    table.cell(1, 0).text = '001'; table.cell(1, 1).text = '<img src=x onerror=alert(1)>'
    doc.add_paragraph('[after](https://invalid.example)')
    data = io.BytesIO(); doc.save(data)
    selected = confirm(repo, filename='structured.docx', content=data.getvalue())
    app = open_page(monkeypatch, repo)
    assert not app.error and len(app.dataframe) == 2
    assert app.dataframe[1].value.iloc[1]['列 1'] == '001'
    assert app.dataframe[1].value.iloc[1]['列 2'] == '<img src=x onerror=alert(1)>'
    assert any(item.value == '<script>alert("before")</script>' for item in app.text)
    assert any(item.value == '[after](https://invalid.example)' for item in app.text)
    assert not app.get('html')
    assert app.text_area(key=f"knowledge_read_{selected['version_id']}").value == selected['text']


def test_search_optional_allowed_types_forwarded_with_same_principal(repo, monkeypatch):
    import enterprise.knowledge_release as release
    calls = []
    def search(query, **kwargs):
        calls.append((query, kwargs))
        return [], {'retrieval_mode': 'fixture', 'reason': 'no_fixture_match'}
    monkeypatch.setattr(release, 'get_search_engine', lambda **kwargs: SimpleNamespace(search=search))
    app = open_page(monkeypatch, repo)
    assert set(app.multiselect(key='knowledge_query_types').options) == {
        '产品配方', '生产工艺', '设备参考', '法规原文', '法规摘要', '行业基准', '市场行情', '成本观察基线', '其他 / 未识别类型'}
    app.text_input(key='knowledge_query').set_value('测试查询')
    app.button(key='knowledge_search').click().run()
    assert not app.exception and calls[-1][1]['knowledge_types'] is None
    app.multiselect(key='knowledge_query_types').set_value(['process', 'market_prices'])
    app.button(key='knowledge_search').click().run()
    assert not app.exception and calls[-1][1]['knowledge_types'] == ['process', 'market_prices']
    assert calls[-1][0] == '测试查询' and calls[-1][1]['principal'] == READER
    assert set(calls[-1][1]['knowledge_types']) <= KNOWLEDGE_TYPES
    assert not repo.db_path.exists(), 'read/search fixture must not initialize any repository'


@pytest.mark.parametrize('role', ['context_only', 'document_basis', 'benchmark_reference', 'market_reference', 'observed_baseline'])
def test_roles_preserve_authority_claim_boundary_and_explicit_role(helpers, role):
    current = {'category': '生产工艺', 'business_metadata': {'authority': 'competition_reference',
               'evidence_role': 'context_only', 'claim_boundary': '原有冲突必须保留', 'known_conflicts': ['示例冲突']}}
    metadata = helpers._registration_metadata(current, role, '生产工艺', {'format': 'txt'}, '10g')
    assert metadata['evidence_role'] == role and metadata['authority'] == 'competition_reference'
    assert '原有冲突必须保留' in metadata['claim_boundary']
    assert helpers.ROLE_BOUNDARIES[role] in metadata['claim_boundary']
    assert metadata['known_conflicts'] == ['示例冲突'] and metadata['specification'] == '10g'
    assert current['business_metadata']['evidence_role'] == 'context_only'
    assert helpers._registration_metadata(None, role, '其他', {'format': 'txt'}, '')['authority'] == 'unreviewed'


@pytest.mark.parametrize('category,role', [('行业基准', 'benchmark_reference'), ('市场参考', 'market_reference'),
                                          ('派生成本基线', 'observed_baseline'), ('异常处理记录', 'context_only')])
def test_reference_cannot_be_promoted_by_category_or_previous_role(helpers, category, role):
    with pytest.raises(ValueError, match='不能登记为机制依据'):
        helpers._registration_metadata(None, 'document_basis', category, {'format': 'txt'}, '')
    current = {'category': category, 'business_metadata': {'evidence_role': role}}
    with pytest.raises(ValueError, match='不能登记为机制依据'):
        helpers._registration_metadata(current, 'document_basis', '生产工艺', {'format': 'txt'}, '')


def test_csv_schema_is_reference_even_when_display_category_is_mechanism(helpers):
    parsed = {'format': 'csv', 'metadata': {'preview_rows': [list(MARKET_HEADERS)]}}
    with pytest.raises(ValueError, match='不能登记为机制依据'):
        helpers._registration_metadata(None, 'document_basis', '生产工艺', parsed, '')


def test_registration_form_has_actual_categories_and_retains_reference_role(repo, monkeypatch):
    selected = confirm(repo, title='对标原件', filename='benchmark.txt', text='对标参考旧版', category='行业基准',
                       metadata={'evidence_role': 'benchmark_reference', 'authority': 'industry_reference',
                                 'claim_boundary': '区间不能视为确定节约'})
    content = b'Updated reference only'
    monkeypatch.setattr(st, 'file_uploader', lambda *args, **kwargs:
                        SimpleNamespace(name='benchmark.txt', size=len(content), getvalue=lambda: content))
    app = open_page(monkeypatch, repo, ADMIN)
    app.selectbox(key='knowledge_target').set_value(selected['doc_id']).run()
    assert not app.exception
    categories = next(item for item in app.selectbox if item.label == '资料类别')
    # 下拉选项以「板块 · 类别」展示，此处还原原始类别值再断言
    raw_categories = {value.split(' · ', 1)[-1] for value in categories.options}
    assert {'法规原文', '法规摘要', '产品配方', '生产工艺', '设备参考', '行业基准', '市场参考', '派生成本基线', '异常处理记录'} <= raw_categories
    purpose = next(item for item in app.selectbox if item.label == '依据用途')
    assert purpose.value == 'benchmark_reference' and len(purpose.options) == 5
    next(item for item in app.button if item.label == '生成差异预览').click().run()
    assert not app.exception and not app.error
    metadata = app.session_state['knowledge_stage']['business_metadata']
    assert metadata['authority'] == 'industry_reference' and metadata['evidence_role'] == 'benchmark_reference'
    assert '区间不能视为确定节约' in metadata['claim_boundary']
    assert len(repo.history()) == 1, 'preview must not confirm a version'
    next(item for item in app.selectbox if item.label == '依据用途').set_value('document_basis')
    next(item for item in app.button if item.label == '生成差异预览').click().run()
    assert not app.exception and any('不能登记为机制依据' in item.value for item in app.error)
    assert len(repo.history()) == 1


def test_board_mapping_covers_catalog_and_anomaly_category(helpers):
    from enterprise.document_classifier import BOARD_LABELS, CATEGORIES
    for category in CATEGORIES:
        assert helpers.board_of(category) in BOARD_LABELS
    assert helpers.board_of('产品配方') == '产品知识'
    assert helpers.board_of('法规原文') == '行业知识'
    assert helpers.board_of('异常处理记录') == '企业内部知识'
    assert helpers.board_of('未知历史类别') == '未归类'


def test_vision_enhancement_flow_stages_merged_draft(repo, monkeypatch):
    """纯扫描PDF：预览失败 → 视觉增强 → 并入登记文本 → 差异预览 → 确认登记（零网络）。"""
    import hashlib
    from tests.test_vision_enhancement import scanned_pdf_bytes
    content = scanned_pdf_bytes()
    monkeypatch.setattr(st, 'file_uploader', lambda *args, **kwargs:
                        SimpleNamespace(name='扫描件.pdf', size=len(content), getvalue=lambda: content))
    import enterprise.multimodal as multimodal
    monkeypatch.setattr(multimodal, 'provider_config',
                        lambda: ('zhipu', 'glm-4v-plus', 'sk-fixture', 'https://fixture.invalid'))
    import enterprise.vision_enhancement as vision
    monkeypatch.setattr(vision, 'enhance_pdf',
        lambda content_bytes, parsed, task='扫描件页面', **kwargs: {
            'ok': True, 'sections': [{'page': 1, 'task': task, 'text': '视觉解析结果X'}],
            'merged_text': '【视觉增强解析稿】视觉解析结果X',
            'meta': {'provider': 'zhipu', 'model': 'glm-4v-plus'}, 'reason': '', 'pages_missing': []})
    app = open_page(monkeypatch, repo, ADMIN)
    assert not app.exception
    next(item for item in app.button if item.label == '开始视觉增强').click().run()
    assert not app.exception
    next(item for item in app.button if item.label == '将增强稿并入登记文本（原件仍为原PDF）').click().run()
    assert not app.exception
    next(item for item in app.text_input if item.label == '适用产品（逗号分隔；公开资料留空）').set_value('甲产品')
    next(item for item in app.text_input if item.label == '适用工厂（逗号分隔；公开资料留空）').set_value('一厂')
    next(item for item in app.button if item.label == '生成差异预览').click().run()
    assert not app.exception
    # 页面仍显示"文本层缺失"预览提示属预期（扫描件）；增强稿应已作为正文进入暂存
    staged = app.session_state['knowledge_stage']
    assert not staged['errors'] and '视觉解析结果X' in staged['text']
    next(item for item in app.checkbox if '已核对正文' in item.label).set_value(True)
    next(item for item in app.text_input if item.label == '登记 / 更新原因').set_value('视觉增强核准')
    next(item for item in app.button if item.label == '确认登记此版本').click().run()
    assert not app.exception
    versions = repo.history()
    assert len(versions) == 1
    assert '视觉解析结果X' in versions[0]['text']
    assert versions[0]['parser'] == 'vision_enhanced_pypdf'
    # 原件仍是原 PDF 字节
    assert versions[0]['sha256'] == hashlib.sha256(content).hexdigest()


def test_sidebar_native_geometry_not_forced():
    from app_pages.design import STYLE
    assert 'stSidebar' not in STYLE
    assert '248px' not in STYLE
    assert 'min-width:248px!important' not in STYLE
