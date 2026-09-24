"""Synthetic-only ingestion coverage, stable format UI and bitemporal update tests.

No real originals/databases, OCR, model loading, server or network access.
"""
from __future__ import annotations

import ast
from dataclasses import replace
import hashlib
import io
from pathlib import Path
import socket
from types import SimpleNamespace

import pytest

from enterprise.knowledge import Repository, preview_file
from enterprise.knowledge_release import ReleaseRepository, get_search_engine
from enterprise.security import Principal

ROOT = Path(__file__).resolve().parents[1]
PAGE = ROOT / 'app_pages/knowledge.py'
ADMIN = Principal('synthetic-ingestion-admin', 'SYNTHETIC_FIXTURE', ('knowledge_admin',), ('*',), ('*',))
READER = replace(ADMIN, user_id='synthetic-reader', roles=('analyst',), factories=('测试厂',), products=('测试品',))


@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    def denied(*args, **kwargs):
        raise AssertionError('Synthetic ingestion tests must not use network')
    monkeypatch.setattr(socket.socket, 'connect', denied)
    monkeypatch.setattr(socket.socket, 'connect_ex', denied)
    monkeypatch.setattr(socket, 'create_connection', denied)


@pytest.fixture
def repo(tmp_path):
    return Repository(tmp_path / 'synthetic-managed', principal=ADMIN)


@pytest.fixture(scope='module')
def helpers():
    tree = ast.parse(PAGE.read_text(encoding='utf-8'), filename=str(PAGE))
    body = []
    for node in tree.body:
        if isinstance(node, ast.Assign) and any(isinstance(target, ast.Tuple) for target in node.targets):
            break
        body.append(node)
    namespace = {}
    exec(compile(ast.Module(body=body, type_ignores=[]), str(PAGE), 'exec'), namespace)
    return SimpleNamespace(**namespace)


def docx_bytes(text, *, drawing=False, external=False):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.opc.constants import RELATIONSHIP_TYPE
    document = Document()
    paragraph = document.add_paragraph(text)
    if drawing:
        paragraph.add_run()._r.append(OxmlElement('w:drawing'))
    if external:
        document.part.relate_to('https://invalid.example/never-fetch.png', RELATIONSHIP_TYPE.IMAGE, is_external=True)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def pdf_bytes(*, text=True, image=False, inline=False, nested=False, blank_second=False):
    from pypdf import PdfWriter
    from pypdf.generic import (ArrayObject, DecodedStreamObject, DictionaryObject,
                              NameObject, NumberObject)
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                             NameObject('/Subtype'): NameObject('/Type1'),
                             NameObject('/BaseFont'): NameObject('/Helvetica')})
    resources = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
    commands = b'BT /F1 12 Tf 30 700 Td (Synthetic local text) Tj ET\n' if text else b''
    if image:
        picture = DecodedStreamObject()
        picture.set_data(b'\xff\x00\x00')
        picture.update({NameObject('/Type'): NameObject('/XObject'), NameObject('/Subtype'): NameObject('/Image'),
                        NameObject('/Width'): NumberObject(1), NameObject('/Height'): NumberObject(1),
                        NameObject('/ColorSpace'): NameObject('/DeviceRGB'), NameObject('/BitsPerComponent'): NumberObject(8)})
        xobjects = DictionaryObject({NameObject('/Im0'): picture})
        if nested:
            form = DecodedStreamObject()
            form.set_data(b'/Im0 Do')
            form.update({NameObject('/Type'): NameObject('/XObject'), NameObject('/Subtype'): NameObject('/Form'),
                         NameObject('/BBox'): ArrayObject([NumberObject(0), NumberObject(0), NumberObject(1), NumberObject(1)]),
                         NameObject('/Resources'): DictionaryObject({NameObject('/XObject'): xobjects})})
            xobjects = DictionaryObject({NameObject('/Fm0'): form})
        resources[NameObject('/XObject')] = xobjects
        commands += b'/Fm0 Do\n' if nested else b'/Im0 Do\n'
    if inline:
        commands += b'BI /W 1 /H 1 /CS /RGB /BPC 8 ID \xff\x00\x00 EI\n'
    page[NameObject('/Resources')] = resources
    stream = DecodedStreamObject(); stream.set_data(commands)
    page[NameObject('/Contents')] = stream
    if blank_second:
        writer.add_blank_page(width=612, height=792)
    buffer = io.BytesIO(); writer.write(buffer)
    return buffer.getvalue()


@pytest.mark.parametrize('kind', ['xobject', 'inline', 'nested'])
def test_pdf_image_coverage_is_visible_without_ocr(kind):
    content = pdf_bytes(image=kind != 'inline', inline=kind == 'inline', nested=kind == 'nested')
    parsed = preview_file(content, 'synthetic.pdf')
    assert not parsed['errors'], parsed
    assert parsed['metadata']['coverage']['image_pages'] == [1]
    assert parsed['metadata']['coverage']['ocr_performed'] is False
    assert parsed['metadata']['coverage']['manual_review_required'] is True
    assert 'OCR' in ''.join(parsed['metadata']['warnings'])
    assert '人工' in ''.join(parsed['metadata']['warnings'])
    assert parsed['text'] == '[第1页]\nSynthetic local text'
    assert parsed['sha256'] == hashlib.sha256(content).hexdigest()
    assert parsed['text_sha256'] == hashlib.sha256(parsed['text'].encode()).hexdigest()


def test_pdf_missing_text_page_is_not_silently_complete():
    parsed = preview_file(pdf_bytes(blank_second=True), 'mixed.pdf')
    assert not parsed['errors']
    assert parsed['metadata']['coverage']['pages_without_text'] == [2]
    assert parsed['metadata']['coverage']['manual_review_required']
    assert parsed['metadata']['warnings']
    assert '[第2页]' in parsed['text']


def test_pure_scan_rejected_before_any_stage(repo):
    content = pdf_bytes(text=False, image=True)
    parsed = preview_file(content, 'scan.pdf')
    assert parsed['errors'] and 'OCR' in ''.join(parsed['errors'])
    assert parsed['metadata']['coverage']['image_pages'] == [1]
    pending = stage(repo, content, 'pdf')
    assert pending['errors'] and pending['stage_id'] is None
    assert repo.history() == []


def test_docx_drawing_and_external_relationship_never_fetched():
    content = docx_bytes('<script>alert(1)</script> [link](https://invalid.example)', drawing=True, external=True)
    parsed = preview_file(content, 'drawing.docx')
    assert not parsed['errors'], parsed
    assert parsed['metadata']['coverage']['drawing_count'] == 1
    assert parsed['metadata']['coverage']['ocr_performed'] is False
    assert parsed['metadata']['coverage']['manual_review_required']
    assert 'OCR' in ''.join(parsed['metadata']['warnings'])
    assert parsed['text'].startswith('<script>')


def test_text_only_docx_and_txt_keep_parser_hash_contract():
    for ext, content in [('docx', docx_bytes('测试品 提取工艺')), ('txt', '测试品 提取工艺'.encode())]:
        parsed = preview_file(content, 'plain.' + ext)
        assert not parsed['errors'] and parsed['text'] == '测试品 提取工艺'
        assert parsed['sha256'] == hashlib.sha256(content).hexdigest()
        assert parsed['text_sha256'] == hashlib.sha256('测试品 提取工艺'.encode()).hexdigest()
        if ext == 'docx':
            assert parsed['metadata']['warnings'] == []
            assert not parsed['metadata']['coverage']['manual_review_required']


def test_legacy_doc_still_explicitly_rejected():
    parsed = preview_file(b'not-docx', 'legacy.doc')
    assert '.docx' in ''.join(parsed['errors']) and parsed['text'] == ''


def stage(repo, content, ext, *, doc_id=None, effective_from='2026-01-01'):
    return repo.stage(content, 'synthetic.' + ext, '合成工艺文档', ['测试品'], effective_from,
                      '生产工艺', ADMIN.user_id, doc_id=doc_id, scope_factories=['测试厂'],
                      visibility='scoped', metadata={'evidence_role': 'document_basis', 'fixture': 'SYNTHETIC_FIXTURE'})


@pytest.mark.parametrize('ext', ['txt', 'docx'])
def test_explicit_two_version_update_and_old_period_replay(repo, ext, monkeypatch):
    import enterprise.knowledge_release as release_module
    monkeypatch.setattr(release_module, '_model_handle', lambda *a, **k: pytest.fail('No model loading'))
    def payload(text):
        return text.encode() if ext == 'txt' else docx_bytes(text)
    original_bytes = payload('测试品 提取工艺 旧版核对')
    first_pending = stage(repo, original_bytes, ext)
    assert not first_pending['errors']
    assert repo.effective_versions('测试品', '2026-06-30', factory='测试厂') == []
    first = repo.commit(first_pending['stage_id'], ADMIN.user_id, '合成旧版确认')
    releases = ReleaseRepository(repository=repo)
    first_release = releases.publish(principal=ADMIN, embedding_model_path='')
    assert first_release['status'] == 'published'
    engine = get_search_engine(repository=repo, embedding_model_path='')
    def search(day):
        return engine.search('测试品 提取工艺', principal=READER, product='测试品', factory='测试厂', as_of=day)[0]
    assert {row['version_id'] for row in search('2026-07-01')} == {first['version_id']}
    original = repo.get(version_id=first['version_id'])
    pending = stage(repo, payload('测试品 提取工艺 新版核对'), ext, doc_id=first['doc_id'], effective_from='2026-07-01')
    assert not pending['errors'] and pending['change']['kind'] == 'new_version'
    assert {row['version_id'] for row in search('2026-07-01')} == {first['version_id']}
    second = repo.commit(pending['stage_id'], ADMIN.user_id, '合成新版确认')
    assert all(row['version_id'] != second['version_id'] for row in search('2026-07-01'))
    second_release = releases.publish(principal=ADMIN, embedding_model_path='')
    assert second_release['status'] == 'published'
    assert {row['version_id'] for row in search('2026-07-01')} == {second['version_id']}
    assert {row['version_id'] for row in search('2026-06-30')} == {first['version_id']}
    assert repo.get(version_id=first['version_id']) == original
    assert repo.read_blob(first['sha256'], version_id=first['version_id']) == original_bytes


def test_empty_authorized_library_keeps_all_supported_formats(repo, monkeypatch):
    from streamlit.testing.v1 import AppTest
    from enterprise.application import Application
    import app_pages._shared as shared
    monkeypatch.setattr(shared, 'streamlit_principal', lambda: READER)
    monkeypatch.setattr(shared, 'Application', lambda actor: Application(actor, repo.root))
    script = 'import streamlit as st\n' + f'st.navigation([st.Page({str(PAGE)!r})]).run()\n'
    app = AppTest.from_string(script, default_timeout=30).run()
    assert not app.exception and not app.error
    assert app.selectbox(key='knowledge_browse_format').options == ['全部授权格式', 'PDF', 'Word(DOCX)', 'TXT', 'CSV']
    for ext in ('pdf', 'docx', 'txt', 'csv'):
        app.selectbox(key='knowledge_browse_format').set_value(ext).run()
        assert not app.exception and any('当前授权范围' in item.value and '没有' in item.value for item in app.info)
    assert any('.doc' in item.value and '.docx' in item.value for item in app.caption)
    assert repo.history() == []
    assert not repo.root.exists()


def test_coverage_warning_is_literal_not_markdown_html_or_remote(helpers, monkeypatch):
    captured = []
    fake = SimpleNamespace(caption=lambda value: captured.append(('caption', value)),
                           warning=lambda value: captured.append(('warning', value)),
                           text=lambda value: captured.append(('text', value)),
                           info=lambda value: captured.append(('info', value)))
    monkeypatch.setitem(helpers._render_parse_coverage.__globals__, 'st', fake)
    unsafe = '<script>alert(1)</script> ![remote](https://invalid.example/image.png)'
    helpers._render_parse_coverage({'coverage': {}, 'warnings': [unsafe]}, 'docx')
    assert ('text', unsafe) in captured
    assert all(unsafe not in value for kind, value in captured if kind != 'text')
    helpers._render_parse_coverage({}, 'pdf')
    assert any(kind == 'info' and '未重解析' in value for kind, value in captured)
