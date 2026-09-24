"""Offline native reading exports from real temporary manufacturing service runs.

Five labelled SIMULATION bundles; no paid models, production DB, delivery writes,
retrieval workaround or source replacement. Exports are written only to tmp_path.
"""
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import re
import socket
import zipfile

from docx import Document
from lxml import etree
from pypdf import PdfReader
import pytest

from enterprise.manufacturing_report import (
    ManufacturingReportError, export_manufacturing_report, validate_manufacturing_run,
)
from enterprise.manufacturing_repository import ManufacturingRepository, digest
from enterprise.manufacturing_service import ManufacturingService
from enterprise.security import Principal

ROOT = Path(__file__).resolve().parents[1]
INDUSTRIES = ('pharma', 'machinery', 'auto_parts', 'chemicals', 'electronics')
ADMIN = Principal('manufacturing-export-fixture', 'SIMULATION exporter', ('system_admin', 'supervisor'), ('*',), ('*',))
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}


def seal(value):
    value['analysis_hash'] = digest({key: item for key, item in value.items() if key not in ('analysis_run_id', 'analysis_hash')})
    return value


def deny(*args, **kwargs):
    raise AssertionError('Frozen exporter must not call live models, retrieval, facts or database')


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    monkeypatch.setenv('COST_TENANT_ID', 'default')
    monkeypatch.setattr(socket.socket, 'connect', deny)
    monkeypatch.setattr('attribution_runtime.run_stage', deny)


@pytest.fixture(scope='module', params=INDUSTRIES)
def frozen_run(request, tmp_path_factory):
    industry = request.param
    folder = ROOT / 'config' / 'manufacturing_examples' / industry
    profile = json.loads((folder / 'domain.json').read_text(encoding='utf-8-sig'))
    adapter = json.loads((folder / 'adapter.json').read_text(encoding='utf-8-sig'))
    files = {family: (family + '.csv', (folder / (family + '.csv')).read_bytes())
             for family in ('actual', 'budget', 'materials', 'labor', 'overhead')}
    root = tmp_path_factory.mktemp('manufacturing-export-' + industry)
    with pytest.MonkeyPatch.context() as patch:
        patch.setenv('COST_TENANT_ID', 'default')
        patch.setattr(socket.socket, 'connect', deny)
        patch.setattr('attribution_runtime.run_stage', deny)
        repository = ManufacturingRepository(root, ADMIN)
        repository.install_profile(profile, adapter, expected_version=0, reason='SIMULATION isolated export test')
        stage = repository.stage(profile['id'], files, periods=['2026-05', '2026-06'], expected_revision=0)
        repository.confirm(stage['stage_id'], expected_revision=0, reason='SIMULATION fixture reviewed')
        service = ManufacturingService(root, ADMIN)
        product = profile['products'][0]
        result = service.analyze(profile['id'], product=product['name'], specification=product['specification'],
                                 month='2026-06', use_llm=False)
        result = service.get_run(result['analysis_run_id'])
    assert result['retrieval_diagnostics']['degraded'] is True
    assert result['retrieval_diagnostics']['formal_model_blocked'] is True
    assert all(not branch['used_llm'] and branch['model_run']['provider_call_count'] == 0
               for branch in result['analyses'].values())
    yield industry, root, result
    assert not (root / 'task_workflow.db').exists()


def docx_text(raw):
    document = Document(BytesIO(raw))
    return '\n'.join([*[p.text for p in document.paragraphs],
                      *[cell.text for table in document.tables for row in table.rows for cell in row.cells]])


def pdf_text(raw):
    return '\n'.join(page.extract_text() for page in PdfReader(BytesIO(raw)).pages)


def compact(text):
    return ''.join(text.split())


def compact_without_pdf_footers(text):
    # Strip only complete footer lines before losing whitespace boundaries; a
    # page-6 footer followed by 50.00% must not become a greedy match for page 650.
    return compact(re.sub(r'(?m)^[ \t]*未签发阅读副本[ \t]*·[ \t]*[0-9]+[ \t]*\r?$', '', text))


@pytest.mark.parametrize('newline', ['\n', '\r\n'])
def test_pdf_footer_normalization_preserves_next_line_digits(newline):
    text = newline.join(('占全量明细', '未签发阅读副本 · 6', '50.00%）'))
    assert compact_without_pdf_footers(text) == '占全量明细50.00%）'


@pytest.mark.parametrize('text', [
    '正文提及未签发阅读副本 · 6，不是页脚。',
    '未签发阅读副本 · 6 后接正文。',
    '正文前缀：未签发阅读副本 · 6',
    '未签发阅读副本 · 6.50',
    '未签发阅读副本',
])
def test_pdf_footer_normalization_preserves_prose_mentions(text):
    assert compact_without_pdf_footers(text) == compact(text)


@pytest.mark.parametrize('format', ['docx', 'pdf'])
def test_five_actual_service_snapshots_export_native_reading_copies(frozen_run, tmp_path, format):
    industry, root, frozen = frozen_run
    before = deepcopy(frozen)
    raw = export_manufacturing_report(frozen, format)
    (tmp_path / (industry + '.' + format)).write_bytes(raw)
    assert raw.startswith(b'PK' if format == 'docx' else b'%PDF')
    text = docx_text(raw) if format == 'docx' else pdf_text(raw)
    normalized = compact(text)
    searchable = compact_without_pdf_footers(text) if format == 'pdf' else normalized
    for value in (frozen['scope']['product'], frozen['scope']['specification'],
                  frozen['scope']['home_factory'], frozen['scope']['peer_factory'],
                  frozen['scope']['month'], frozen['measurement']['quantity_unit'],
                  frozen['measurement']['unit_cost_unit'], frozen['analysis_hash'],
                  'SIMULATION', 'needs_human_review', '未签发', 'JSON为审计权威原件'):
        assert compact(value) in normalized
    for kind, branch in frozen['analyses'].items():
        for section in branch['narrative']['sections']:
            # Inline source marks become native footnotes/superscripts; the frozen
            # numeric wording and immediate action themselves are not recalculated.
            numeric = section.get('numeric_explanation', '')
            # Native PDF citations remain searchable superscript numbers between
            # the unchanged prose fragments, and page footers may split a paragraph.
            fragments = re.split(r'\[[A-Za-z][A-Za-z0-9_-]*\]|\n', numeric)
            assert all(compact(fragment) in searchable for fragment in fragments if fragment.strip())
            assert compact(section['immediate_action']) in normalized
        criteria = compact(branch['narrative'].get('followup_criteria', ''))
        if criteria:
            assert normalized.count(criteria) == sum(
                compact(row['narrative'].get('followup_criteria', '')) == criteria for row in frozen['analyses'].values())
    for row in (frozen['facts']['current'], frozen['facts']['budget']['baseline'], frozen['facts']['benchmark']['peer']):
        for field in ('total', 'output', 'unitcost'):
            assert row[field] in text
    assert '中药一厂' not in text and '银黄口服液' not in text and '元/盒' not in text
    other_products = [row['name'] for row in frozen['profile']['products'] if row['name'] != frozen['scope']['product']]
    assert not any(compact(name) in normalized for name in other_products)
    assert '2026-07' not in text
    assert frozen == before
    assert not (root / 'task_workflow.db').exists()


def test_docx_native_headings_tables_footnotes_and_suggestion_only_border(frozen_run):
    _, _, frozen = frozen_run
    raw = export_manufacturing_report(frozen, 'docx')
    document = Document(BytesIO(raw))
    assert document.tables and any(p.style.name == 'Heading 1' for p in document.paragraphs)
    with zipfile.ZipFile(BytesIO(raw)) as archive:
        xml = etree.fromstring(archive.read('word/document.xml'))
        assert 'word/footnotes.xml' in archive.namelist()
        notes = etree.fromstring(archive.read('word/footnotes.xml'))
        assert notes.xpath('//w:footnote[@w:id="1"]', namespaces=NS)
        assert xml.xpath('//w:footnoteReference', namespaces=NS)
        assert xml.xpath('//w:tblHeader', namespaces=NS)
        border_paragraphs = xml.xpath('//w:p[w:pPr/w:pBdr/w:left]', namespaces=NS)
        assert border_paragraphs
        for paragraph in border_paragraphs:
            assert paragraph.xpath('./w:pPr/w:pBdr/w:left/@w:color', namespaces=NS) == ['2A78D6']
            assert paragraph.xpath('./w:pPr/w:pBdr/w:left/@w:sz', namespaces=NS) == ['12']
            prefix = paragraph.xpath('./w:r[w:rPr/w:b and w:rPr/w:color[@w:val="2A78D6"]]/w:t/text()', namespaces=NS)
            assert prefix == ['建议']
        style = etree.fromstring(archive.read('word/styles.xml'))
        for name in ('Heading1', 'Heading2', 'Heading3'):
            assert style.xpath(f'//w:style[@w:styleId="{name}"]/w:rPr/w:color/@w:val', namespaces=NS) == ['0B0B0B']
    assert document.core_properties.subject == '冻结JSON ' + frozen['analysis_hash'] + '；未签发阅读副本'


def test_pdf_searchable_embedded_font_blue_advice_and_page_fit(frozen_run):
    _, _, frozen = frozen_run
    raw = export_manufacturing_report(frozen, 'pdf')
    reader = PdfReader(BytesIO(raw))
    streams = b'\n'.join(page.get_contents().get_data() for page in reader.pages)
    # The left border is drawn as a real PDF line in the requested RGB color;
    # no image-only page or simulated text bar stands in for the advice block.
    assert b'.164706 .470588 .839216 RG' in streams
    assert b'1.5 w' in streams
    embedded = []
    for page in reader.pages:
        assert float(page.mediabox.width) == pytest.approx(595.2756, abs=0.1)
        assert float(page.mediabox.height) == pytest.approx(841.8898, abs=0.1)
        for item in page['/Resources'].get('/Font', {}).values():
            font = item.get_object()
            descriptor = font.get('/FontDescriptor')
            if descriptor:
                embedded.append('/FontFile2' in descriptor.get_object())
    assert embedded and all(embedded)
    # Optional installed visual parser verifies no body span extends outside the
    # A4 page. It never creates or edits production artifacts.
    fitz = pytest.importorskip('fitz')
    pdf = fitz.open(stream=raw, filetype='pdf')
    for page in pdf:
        for block in page.get_text('dict')['blocks']:
            for line in block.get('lines', []):
                for span in line.get('spans', []):
                    left, top, right, bottom = span['bbox']
                    assert left >= 0 and right <= page.rect.width + 0.5
                    assert top >= 0 and bottom <= page.rect.height + 0.5


def test_export_never_reloads_live_inputs_or_uses_source_paths(frozen_run, monkeypatch):
    _, _, frozen = frozen_run
    value = deepcopy(frozen)
    for branch in value['analyses'].values():
        source = branch['sources'][0]
        if source['source'].get('records'):
            source['source']['records'][0]['file'] = r'Z:\DO-NOT-OPEN\secrets.csv'
        else:
            source['source']['file'] = r'Z:\DO-NOT-OPEN\secrets.csv'
    seal(value)
    monkeypatch.setattr('enterprise.manufacturing_runtime.ManufacturingRuntime.analysis', deny)
    monkeypatch.setattr('enterprise.manufacturing_projection.project_manufacturing_analysis', deny)
    monkeypatch.setattr('enterprise.manufacturing_repository.ManufacturingRepository.current', deny)
    monkeypatch.setattr('enterprise.manufacturing_service.ManufacturingService._connect', deny)
    monkeypatch.setattr('enterprise.analysis_service.report_evidence', deny)
    monkeypatch.setattr('enterprise.analysis_narrative.build_attribution_narrative', deny)
    monkeypatch.setattr('enterprise.analysis_narrative.build_benchmark_narrative', deny)
    monkeypatch.setattr('report.model.load_data', deny)
    for format in ('docx', 'pdf'):
        raw = export_manufacturing_report(value, format)
        assert raw
        text = docx_text(raw) if format == 'docx' else pdf_text(raw)
        assert compact(r'Z:\DO-NOT-OPEN\secrets.csv') in compact(text)


def test_export_preserves_exact_adopted_advice_and_raw_quote_offsets(frozen_run):
    _, _, original = frozen_run
    value = deepcopy(original)
    branch = value['analyses']['attribution']
    section = branch['narrative']['sections'][0]
    # Explicit export-only synthetic frozen value, not evidence of real LLM use.
    accepted = 'TEST_ONLY exact accepted advice：采购负责人核对结算单；保留未闭合项，不确认节约。'
    quote = 'TEST_ONLY 原文：the controlled instruction requires reconciled material issue records.'
    source = {'id': 'K_EXPORT_FIXTURE', 'kind': 'document_basis', 'elements': ['材料'],
              'text': quote, 'support_status': 'eligible',
              'scope': {'product': value['scope']['product'], 'specification': value['scope']['specification'],
                        'months': [value['scope']['month']]},
              'source': {'file': 'TEST_ONLY-reference.txt', 'offset': 37, 'end_offset': 37 + len(quote),
                         'page': 8, 'line': 4, 'sha256': 'a' * 64, 'quote': quote}}
    branch['sources'].append(source)
    branch['narrative']['evidence_ids'].append(source['id'])
    section['evidence_ids'].append(source['id'])
    section['accepted_model_recommendation'] = accepted
    section['mechanism_note'] = quote + ' [K_EXPORT_FIXTURE]'
    seal(value)
    for format in ('docx', 'pdf'):
        raw = export_manufacturing_report(value, format)
        text = docx_text(raw) if format == 'docx' else pdf_text(raw)
        assert compact(accepted) in compact(text)
        assert compact(quote) in compact(text)
        assert 'offset 37' in text and '来源标注页 8' in text
        assert 'K_EXPORT_FIXTURE' in text
        assert 'source_row_sha256' in text and 'source_fields' in text
    assert original != value


def test_bound_prose_exactly_once_replaces_old_main_numeric_paragraph(frozen_run):
    _, _, original = frozen_run
    value = deepcopy(original)
    section = value['analyses']['attribution']['narrative']['sections'][0]
    prose = 'TEST_ONLY_BOUND_PROSE 原样保留的散文首段。\n原样保留的第二段，结论仍待凭证核对。'
    section['prose'] = prose
    section['prose_mode'] = 'bound-numeric-prose/1'
    section['numeric_explanation'] = 'TEST_ONLY_OLD_NUMERIC_BODY_NOT_DISPLAYED'
    section.setdefault('model_core', {})['hypothesis'] = 'TEST_ONLY_OLD_HYPOTHESIS_NOT_DISPLAYED'
    seal(value)
    for format in ('docx', 'pdf'):
        raw = export_manufacturing_report(value, format)
        text = docx_text(raw) if format == 'docx' else pdf_text(raw)
        assert compact(text).count(compact(prose)) == 1
        assert 'TEST_ONLY_OLD_NUMERIC_BODY_NOT_DISPLAYED' not in text
        assert 'TEST_ONLY_OLD_HYPOTHESIS_NOT_DISPLAYED' not in text
        assert compact(section['immediate_action']) in compact(text)
        assert value['facts']['current']['total'] in text
    fallback = deepcopy(value)
    section = fallback['analyses']['attribution']['narrative']['sections'][0]
    section['prose_mode'] = 'unknown-contract'
    seal(fallback)
    text = docx_text(export_manufacturing_report(fallback, 'docx'))
    assert 'TEST_ONLY_OLD_NUMERIC_BODY_NOT_DISPLAYED' in text
    assert 'TEST_ONLY_BOUND_PROSE' not in text


@pytest.mark.parametrize('fault', ['hash', 'schema', 'profile', 'unit', 'scope', 'future', 'reference'])
def test_tampering_rejected_even_if_semantic_change_resealed(frozen_run, fault):
    _, _, original = frozen_run
    value = deepcopy(original)
    if fault == 'hash':
        value['analyses']['attribution']['narrative']['overview'] = 'tampered'
    elif fault == 'schema':
        value['schema_version'] = 'manufacturing-analysis/999'
        seal(value)
    elif fault == 'profile':
        value['profile']['label'] += ' changed'
        seal(value)
    elif fault == 'unit':
        value['measurement']['quantity_unit'] = 'invented-unit'
        seal(value)
    elif fault == 'scope':
        value['facts']['current']['factory'] = 'OTHER FACTORY'
        seal(value)
    elif fault == 'future':
        value['facts']['details']['materials'][0]['current']['month'] = '2026-07'
        seal(value)
    else:
        value['analyses']['attribution']['narrative']['sections'][0]['evidence_ids'].append('MISSING_REF')
        seal(value)
    with pytest.raises(ManufacturingReportError):
        export_manufacturing_report(value, 'docx')
    assert validate_manufacturing_run(original) == original


def test_regular_only_approved_font_still_has_searchable_bold_style(frozen_run, monkeypatch):
    _, _, frozen = frozen_run
    import enterprise.manufacturing_report as report
    original = report._font
    def regular_only(blocks):
        descriptor = original(blocks)
        descriptor.pop('bold_path', None)
        descriptor.pop('bold_sha256', None)
        return descriptor
    monkeypatch.setattr(report, '_font', regular_only)
    raw = export_manufacturing_report(frozen, 'pdf')
    text = pdf_text(raw)
    assert compact(frozen['scope']['product']) in compact(text)
    assert compact(frozen['analyses']['attribution']['narrative']['sections'][0]['immediate_action']) in compact(text)
    streams = b'\n'.join(page.get_contents().get_data() for page in PdfReader(BytesIO(raw)).pages)
    assert b'2 Tr' in streams


def test_unknown_format_and_missing_font_fail_explicitly(frozen_run, monkeypatch):
    _, _, frozen = frozen_run
    for format in ('html', 'doc', 'DOCX', '../../x', None):
        with pytest.raises(ManufacturingReportError):
            export_manufacturing_report(frozen, format)
    def missing(*args, **kwargs):
        raise RuntimeError('test font unavailable')
    monkeypatch.setattr('report.export.font_descriptor', missing)
    with pytest.raises(RuntimeError, match='字体'):
        export_manufacturing_report(frozen, 'pdf')
    with pytest.raises(RuntimeError, match='字体'):
        export_manufacturing_report(frozen, 'docx')
