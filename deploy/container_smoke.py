#!/usr/bin/env python3
"""Offline, synthetic smoke for the built Linux core image, running as UID 10001.

Invoke via Docker with --network none and --entrypoint python. This deliberately
bypasses the production OIDC supervisor only in this disposable test process; it
never starts a listener, worker or mock service. Inputs and generated documents
are synthetic and temporary. No competition template, model or secret is read.
The stdout summary is suitable for CI logs; no artifact is published.
"""
from __future__ import annotations

from contextlib import contextmanager
import hashlib
import importlib
import importlib.util
import io
import json
import os
from pathlib import Path
import socket
import subprocess
import sys
import tempfile

APP_ROOT = Path(__file__).resolve().parents[1]
SAMPLE_TEXT = '合成字体检查：中文成本，数量单位，符号 ±×÷；负号 \u2212；部首 \u2ee9。'
CORE_MODULES = (
    'backend_api', 'enterprise.application', 'enterprise.security',
    'enterprise.knowledge', 'enterprise.knowledge_release', 'enterprise.knowledge_langchain',
    'enterprise.knowledge_context', 'enterprise.analysis_service', 'enterprise.model_gateway',
    'enterprise.operations', 'enterprise.task_workflow', 'enterprise.task_worker',
    'dashboard.data_layer', 'report.model', 'report.export', 'report.registry',
    'attribution_gen', 'attribution_runtime',
)


class SmokeFailure(RuntimeError):
    pass


def require(condition, code):
    if not condition:
        raise SmokeFailure(code)


def digest(raw):
    return hashlib.sha256(raw).hexdigest()


def configure_isolated_environment(root):
    """Set every path before application imports; change this process only."""
    data = root / 'empty-data'
    data.mkdir()
    (root / 'cache').mkdir()
    values = {
        'COST_AUTH_MODE': 'local', 'COST_BIND_HOST': '127.0.0.1',
        'COST_TENANT_ID': 'container-smoke', 'STREAMLIT_SERVER_ADDRESS': '127.0.0.1',
        'STREAMLIT_BROWSER_GATHER_USAGE_STATS': 'false',
        'COST_DATA_DIR': str(data), 'COST_MANAGED_DIR': str(root / 'managed'),
        'COST_CHROMA_PATH': str(root / 'unused-chroma'),
        'COST_LLM_CONFIG_FILE': str(root / 'absent-model-config.json'),
        'COST_LLM_API_KEY': '', 'DASHSCOPE_API_KEY': '', 'OPENAI_API_KEY': '',
        'COST_LLM_APPROVED_CLOUD': 'false', 'COST_WORKER_SUBJECT': '',
        'COST_OFFICIAL_MOCK_SCRIPT': '', 'COST_AUTHORIZATION_FILE': '',
        'BGE_M3_PATH': str(root / 'absent-bge'), 'BGE_RERANKER_PATH': str(root / 'absent-reranker'),
        'HF_HOME': str(root / 'cache' / 'huggingface'), 'HF_HUB_OFFLINE': '1',
        'HF_DATASETS_OFFLINE': '1', 'TRANSFORMERS_OFFLINE': '1',
        'HF_HUB_DISABLE_TELEMETRY': '1', 'DO_NOT_TRACK': '1',
        'LANGCHAIN_TRACING': 'false', 'LANGCHAIN_TRACING_V2': 'false',
        'LANGCHAIN_HANDLER': '', 'LANGCHAIN_API_KEY': '',
        'LANGSMITH_TRACING': 'false', 'LANGSMITH_TRACING_V2': 'false',
        'LANGSMITH_API_KEY': '', 'LANGSMITH_ENDPOINT': 'https://tracing-disabled.invalid',
        'MPLBACKEND': 'Agg', 'MPLCONFIGDIR': str(root / 'cache' / 'matplotlib'),
        'OMP_NUM_THREADS': '1', 'OPENBLAS_NUM_THREADS': '1', 'MKL_NUM_THREADS': '1',
    }
    os.environ.update(values)
    return data


@contextmanager
def reject_socket_io():
    """Reject connections/listeners; count urllib3's bind-only IPv6 probe.

    urllib3 probes IPv6 availability at import by binding an ephemeral ::1
    socket, then closing it without listening or connecting. That local probe
    is recorded separately and is not an outbound request or a service.
    """
    calls = {'connection_attempts': 0, 'listener_attempts': 0,
             'other_bind_attempts': 0, 'loopback_bind_probes': 0}
    saved = {(socket.socket, name): getattr(socket.socket, name)
             for name in ('connect', 'connect_ex', 'bind', 'listen')}
    saved[(socket, 'create_connection')] = socket.create_connection
    def forbidden(kind):
        def reject(*_args, **_kwargs):
            calls[kind] += 1
            raise SmokeFailure('network_or_listener_attempted')
        return reject
    def checked_bind(sock, address):
        if sock.family == socket.AF_INET6 and sock.type == socket.SOCK_STREAM and address == ('::1', 0):
            calls['loopback_bind_probes'] += 1
            return saved[(socket.socket, 'bind')](sock, address)
        calls['other_bind_attempts'] += 1
        raise SmokeFailure('unexpected_bind_attempted')
    try:
        socket.socket.connect = forbidden('connection_attempts')
        socket.socket.connect_ex = forbidden('connection_attempts')
        socket.create_connection = forbidden('connection_attempts')
        socket.socket.listen = forbidden('listener_attempts')
        socket.socket.bind = checked_bind
        yield calls
    finally:
        for (owner, name), original in saved.items():
            setattr(owner, name, original)


def check_imports():
    sys.path.insert(0, str(APP_ROOT))
    for name in CORE_MODULES:
        module = importlib.import_module(name)
        require(Path(module.__file__).resolve().is_relative_to(APP_ROOT), 'core_import_outside_image')
    import streamlit
    from langchain_core.documents import Document
    from langchain_core.retrievers import BaseRetriever
    from enterprise.knowledge_langchain import ControlledKnowledgeRetriever
    require(issubclass(ControlledKnowledgeRetriever, BaseRetriever), 'langchain_adapter_missing')
    require(Document(page_content=SAMPLE_TEXT).page_content == SAMPLE_TEXT, 'langchain_document_invalid')
    require(importlib.util.find_spec('torch') is None, 'unexpected_optional_torch')
    require(importlib.util.find_spec('transformers') is None, 'unexpected_optional_transformers')
    return {'core_modules': len(CORE_MODULES), 'all_inside_image': True,
            'streamlit': streamlit.__version__, 'langchain_base_retriever': True,
            'optional_model_packages': 'absent'}


def check_api(data_root):
    from fastapi.testclient import TestClient
    from backend_api import app
    from paths import DATA_DIR, MANAGED_DIR
    require(DATA_DIR == data_root.resolve() and not any(DATA_DIR.iterdir()), 'data_root_not_empty')
    require(MANAGED_DIR.is_relative_to(data_root.parent), 'managed_root_not_isolated')
    statuses = {}
    with TestClient(app, base_url='http://127.0.0.1', client=('127.0.0.1', 51001)) as client:
        for path in ('/api/health', '/api/me', '/api/system/status', '/api/tasks/summary'):
            response = client.get(path)
            require(response.status_code == 200, 'api_check_failed:' + path)
            require(response.headers.get('cache-control') == 'no-store', 'api_cache_policy_failed')
            body = response.json()
            if path == '/api/health':
                require(body['status'] == 'ok' and len(body['source_fingerprint']) == 64, 'health_identity_failed')
            elif path == '/api/me':
                require(body['auth_method'] == 'local_os_demo', 'unexpected_auth_mode')
            elif path == '/api/system/status':
                require(body['model']['configured'] is False, 'unexpected_model_configuration')
            elif path == '/api/tasks/summary':
                require(body['generated'] == 0 and body['simulated'] is True, 'unexpected_task_state')
            statuses[path] = response.status_code
    with TestClient(app, base_url='http://127.0.0.1', client=('192.0.2.10', 51002)) as client:
        require(client.get('/api/me').status_code == 401, 'non_loopback_identity_not_rejected')
    require(not any(DATA_DIR.iterdir()), 'business_data_created')
    return {'transport': 'in_process_asgi', 'statuses': statuses, 'non_loopback_denied': True,
            'business_data_files': 0, 'tasks_created': 0}


def embedded_font_streams(page):
    streams = []
    for reference in page['/Resources']['/Font'].values():
        font = reference.get_object()
        descendants = font.get('/DescendantFonts', [])
        candidates = [font] + [item.get_object() for item in descendants]
        for candidate in candidates:
            reference = candidate.get('/FontDescriptor')
            if reference is None:
                continue
            descriptor = reference.get_object()
            for key in ('/FontFile', '/FontFile2', '/FontFile3'):
                if key in descriptor:
                    streams.append(descriptor[key].get_object().get_data())
    return streams


def check_document_io(root):
    from docx import Document
    from docx.oxml.ns import qn
    from pypdf import PdfReader
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen.canvas import Canvas
    from report.export import font_descriptor

    selected = font_descriptor(required_text=SAMPLE_TEXT)
    font = TTFont('ContainerSmokeCJK', selected['path'], subfontIndex=selected['subfont_index'])
    require(all(char.isspace() or ord(char) in font.face.charToGlyph for char in SAMPLE_TEXT), 'font_glyph_missing')
    pdfmetrics.registerFont(font)
    pdf_buffer = io.BytesIO()
    canvas = Canvas(pdf_buffer, pagesize=A4, invariant=1)
    canvas.setTitle('Synthetic container smoke')
    canvas.setFont('ContainerSmokeCJK', 11)
    canvas.drawString(36, A4[1] - 54, SAMPLE_TEXT)
    canvas.showPage()
    canvas.save()
    pdf_bytes = pdf_buffer.getvalue()
    pdf_path = root / 'synthetic-smoke.pdf'
    pdf_path.write_bytes(pdf_bytes)
    require(pdf_path.read_bytes() == pdf_bytes, 'pdf_disk_roundtrip_failed')
    reader = PdfReader(io.BytesIO(pdf_path.read_bytes()), strict=True)
    require(len(reader.pages) == 1, 'pdf_page_count_failed')
    extracted = reader.pages[0].extract_text()
    require(all(char.isspace() or char in extracted for char in SAMPLE_TEXT), 'pdf_unicode_roundtrip_failed')
    embedded = embedded_font_streams(reader.pages[0])
    require(any(len(stream) > 100 for stream in embedded), 'pdf_font_not_embedded')

    document = Document()  # python-docx's bundled blank document, no competition template.
    run = document.add_paragraph().add_run(SAMPLE_TEXT)
    run.font.name = selected['family']
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), selected['family'])
    docx_buffer = io.BytesIO()
    document.save(docx_buffer)
    docx_bytes = docx_buffer.getvalue()
    docx_path = root / 'synthetic-smoke.docx'
    docx_path.write_bytes(docx_bytes)
    require(docx_path.read_bytes() == docx_bytes, 'docx_disk_roundtrip_failed')
    reopened = Document(io.BytesIO(docx_path.read_bytes()))
    require(''.join(paragraph.text for paragraph in reopened.paragraphs) == SAMPLE_TEXT, 'docx_unicode_roundtrip_failed')
    return {'synthetic_inputs_only': True, 'competition_template_used': False,
            'font_file': selected['file'], 'font_sha256': selected['sha256'],
            'required_codepoints': ['U+2EE9', 'U+2212'], 'pdf_pages': len(reader.pages),
            'pdf_embedded_font_streams': len(embedded), 'pdf_unicode_roundtrip': True,
            'docx_unicode_roundtrip': True, 'temporary_disk_roundtrip': True,
            'pdf_sha256': digest(pdf_bytes), 'docx_sha256': digest(docx_bytes)}


def main():
    stage = 'platform'
    try:
        require(sys.platform.startswith('linux'), 'linux_container_required')
        require(os.getuid() == 10001 and os.getgid() == 10001, 'uid_gid_10001_required')
        stage = 'pip_check'
        result = subprocess.run([sys.executable, '-m', 'pip', 'check'], stdin=subprocess.DEVNULL,
                                stdout=sys.stdout, stderr=sys.stderr, timeout=90, check=False)
        require(result.returncode == 0, 'pip_check_failed')
        with tempfile.TemporaryDirectory(prefix='core-container-smoke-') as folder:
            root = Path(folder)
            data_root = configure_isolated_environment(root)
            with reject_socket_io() as network:
                stage = 'core_imports'
                imports = check_imports()
                stage = 'api'
                api = check_api(data_root)
                stage = 'font_and_documents'
                documents = check_document_io(root)
                require(not any(network[key] for key in ('connection_attempts', 'listener_attempts',
                                                          'other_bind_attempts')), 'unexpected_socket_attempt')
            report = {'schema_version': 'core-container-smoke/1.0', 'status': 'passed',
                      'uid': os.getuid(), 'gid': os.getgid(), 'pip_check_exit_code': result.returncode,
                      'imports': imports, 'api': api, 'documents': documents,
                      'application_socket_checks': network, 'model_calls': 0, 'mock_calls': 0,
                      'production_oidc_checked': False, 'full_business_report_checked': False,
                      'external_artifacts_published': False}
        print(json.dumps(report, ensure_ascii=True))
        return 0
    except Exception as exc:
        # Static check codes only; no provider replies, data or environment values.
        error = str(exc) if isinstance(exc, SmokeFailure) else type(exc).__name__
        print(json.dumps({'schema_version': 'core-container-smoke/1.0', 'status': 'failed',
                          'stage': stage, 'error': error, 'external_artifacts_published': False}), file=sys.stderr)
        return 1


if __name__ == '__main__':
    raise SystemExit(main())
